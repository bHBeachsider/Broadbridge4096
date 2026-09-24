"""DOCX intake tests use a filled copy of the actual interview guide."""
import copy
from datetime import datetime
import importlib
import json
from pathlib import Path
import subprocess
import sys

from docx import Document
import pytest

PACK = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PACK / "scripts"))
FIXTURE = PACK / "tests/fixtures/Synthetic_Filled_Interview_Guide.docx"


def parse(**kwargs):
    return forms().parse_form(FIXTURE, case_id="FORM-SYN-001", family_id="FORM-FAMILY-001", **kwargs)


def forms():
    return importlib.import_module("case_from_form_fallback")


def parse_path(path, **kwargs):
    return forms().parse_form(path, case_id="FORM-SYN-001", family_id="FORM-FAMILY-001", **kwargs)


def test_fallback_uses_canonical_page_contract_without_alternate_schema():
    case, questions, _ = parse()
    importlib.import_module("case_contract").validate_export({
        "exported_at": "2026-09-24T00:00:00Z", "workflow": {
            "schema": "broadbridge.workflow/1", "updated_at": "", "answers": {}},
        "cases": [case]})
    assert case["schema"] == "broadbridge.case_record/1"
    assert isinstance(questions[0]["evidence_ids"], str)
    assert isinstance(questions[0]["hard_fail_criteria"], str)


def test_downstream_validator_rejects_retired_form_contract():
    case, _, _ = parse()
    case["schema"] = "broadbridge.case_form/1"
    with pytest.raises(ValueError):
        importlib.import_module("case_contract").validate_case(case)


def test_filled_actual_guide_maps_every_b_item_and_part_c():
    case, questions, review = parse()
    assert case["case_id"] == "FORM-SYN-001"
    assert case["status"] == "draft"
    assert case["identity"]["unit_service"] == "Synthetic vacuum wash section"
    assert case["identity"]["record_type"] == "hypothetical"
    assert case["decision_time"]["b1_trigger"] == "Synthetic differential pressure increased; establish which checks are needed."
    assert case["decision_time"]["observations"] == [
        {"time": "T-0", "variable_location": "Bed differential pressure", "value_units_basis": "18 kPa differential", "source": "Synthetic PDT-101", "quality": "unknown"}]
    assert case["hindsight"]["hypotheses"][0]["discriminator"] == "POST-EVENT discriminating inspection"
    assert case["hindsight"]["b13_lesson_and_limits"] == "POST-EVENT lesson limited to this synthetic example."
    assert len(case["evidence"]) == 5
    assert case["evidence"][0]["format"] == "Synthetic drawing"
    assert len(questions) == 3
    assert questions[0]["type"] == "missing_data"
    assert questions[0]["evidence_ids"] == "B1, B4"
    assert questions[0]["hard_fail_criteria"] == "Invents a measurement\nRecommends an unverified operating change"
    assert questions[2]["split"] == "locked_test"
    assert review["workflow"]["A1"] == "Synthetic workflow: collect facts, check instruments, compare mechanisms."
    assert review["consent"]["reviewer_signature"] == "SYNTHETIC REVIEWER ONLY / 2026-09-24"
    assert review["empty_boxes"] == []
    assert case["questions"] == questions
    assert datetime.fromisoformat(case["created_at"]).utcoffset().total_seconds() == 0
    assert case["created_at"] == case["updated_at"]


def test_empty_box_is_empty_string_and_flagged_not_invented(tmp_path):
    document = Document(FIXTURE)
    for table in document.tables:
        if table.cell(0, 0).text.startswith("Synthetic differential pressure"):
            table.cell(0, 0).text = ""
            break
    path = tmp_path / "empty.docx"
    document.save(path)
    case, _, review = parse_path(path)
    assert case["decision_time"]["b1_trigger"] == ""
    assert "decision_time.b1_trigger" in review["empty_boxes"]
    assert case["status"] == "draft"
    importlib.import_module("case_contract").validate_case(case)
    with pytest.raises(ValueError, match="empty"):
        parse_path(path, case_signed_off=True)


def test_explicit_unknown_is_preserved_without_empty_flag(tmp_path):
    document = Document(FIXTURE)
    for table in document.tables:
        if table.cell(0, 0).text.startswith("Synthetic differential pressure"):
            table.cell(0, 0).text = "Unknown"
            break
    path = tmp_path / "unknown.docx"
    document.save(path)
    case, _, review = parse_path(path)
    assert case["decision_time"]["b1_trigger"] == "Unknown"
    assert "decision_time.b1_trigger" not in review["empty_boxes"]


def test_signatures_do_not_automatically_approve_case():
    case, _, _ = parse()
    assert case["reviewer_signoff"]["signed"] is False
    signed, _, _ = parse(case_signed_off=True)
    assert signed["status"] == "signed"
    assert signed["reviewer_signoff"]["signed"] is True
    assert signed["reviewer_signoff"] == {"signed": True, "name": "SYNTHETIC REVIEWER ONLY", "date": "2026-09-24"}


@pytest.mark.parametrize("missing", ["case_id", "signature", "date", "family"])
def test_hand_edited_signed_fallback_cannot_bypass_identity_and_signoff(missing):
    module = importlib.import_module("case_contract")
    case, _, _ = parse(case_signed_off=True)
    for question in case["questions"]: question["split"] = "train"
    if missing == "case_id": case["case_id"] = ""
    if missing == "signature": case["reviewer_signoff"]["name"] = "  "
    if missing == "date": case["reviewer_signoff"]["date"] = ""
    if missing == "family": case["family_id"] = " unknown "
    with pytest.raises(ValueError): module.validate_case(case, training=True)


@pytest.mark.parametrize("ids", [{}, {"case_id": "FORM-SYN-001"}, {"family_id": "F-001"},
    {"case_id": "unknown", "family_id": "F-001"}, {"case_id": "FORM-SYN-001", "family_id": "unknown"}])
def test_reviewed_identifiers_required_even_for_drafts(ids):
    with pytest.raises(ValueError, match="reviewed"):
        forms().parse_form(FIXTURE, **ids)


@pytest.mark.parametrize("label", ["Real event / reconstructed / hypothetical",
    "May this case be used for: training / testing only / reference only", "Type (from table above)",
    "Split: train / dev / locked test"])
@pytest.mark.parametrize("value", ["", "unknown", "undecided"])
def test_missing_or_unknown_enum_requires_explicit_choice(tmp_path, label, value):
    doc = Document(FIXTURE)
    row = next(row for table in doc.tables for row in table.rows if row.cells[0].text == label)
    row.cells[1].text = value
    path = tmp_path / "missing-choice.docx"
    doc.save(path)
    with pytest.raises(ValueError, match="explicit.*choice"):
        parse_path(path)


@pytest.mark.parametrize("change", ["b_missing", "unknown_key", "wrong_enum", "criteria_array", "evidence_ids_array", "observations_object"])
def test_form_schema_rejects_shape_drift(change):
    module = importlib.import_module("case_contract")
    case, _, _ = parse()
    if change == "b_missing": del case["hindsight"]["b13_lesson_and_limits"]
    if change == "unknown_key": case["decision_time"]["hindsight"] = "bad"
    if change == "wrong_enum": case["questions"][0]["type"] = "invented"
    if change == "criteria_array": case["questions"][0]["hard_fail_criteria"] = ["bad"]
    if change == "evidence_ids_array": case["questions"][0]["evidence_ids"] = ["B1"]
    if change == "observations_object": case["decision_time"]["observations"] = {}
    with pytest.raises(ValueError): module.validate_case(case)


def test_prompt_from_form_uses_only_decision_time_and_service():
    brief = importlib.import_module("run_brief")
    case, _, _ = parse()
    messages = brief.run_case(case, dry_run=True)
    payload = json.loads(messages[1]["content"])
    assert set(payload) == {"identity", "decision_time"}
    assert payload["identity"] == {"unit_service": "Synthetic vacuum wash section"}
    assert "POST-EVENT" not in str(messages)
    assert "SCORING ANSWER" not in str(messages)
    assert "b14_evidence" not in str(messages)
    assert "SYNTHETIC REVIEWER" not in str(messages)


def test_unknown_is_not_misclassified_as_retrospective_leak():
    brief = importlib.import_module("run_brief")
    case, _, _ = parse()
    case["decision_time"]["b6_distrusted_or_missing"] = "unknown"
    case["hindsight"]["b11_confidence"] = "unknown"
    case["questions"][0]["reference_answer"] = "unknown"
    brief.run_case(case, dry_run=True)


def test_real_hindsight_leak_in_form_is_blocked():
    brief = importlib.import_module("run_brief")
    case, _, _ = parse()
    case["decision_time"]["b2_operating_context"] = case["hindsight"]["b8_turning_point"]
    with pytest.raises(AssertionError, match="leak"):
        brief.run_case(case, dry_run=True)


def test_fallback_cli_enters_same_import_brief_and_score_flow(tmp_path):
    result = subprocess.run([sys.executable, str(PACK / "scripts/case_from_form_fallback.py"), str(FIXTURE),
        str(tmp_path / "parsed"), "--case-id", "FORM-SYN-001", "--family-id", "FORM-FAMILY-001",
        "--case-signed-off", "--standalone-family"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    output = tmp_path / "parsed"
    case = json.loads((output / "case_record.json").read_text(encoding="utf-8"))
    assert case["status"] == "signed"
    assert len(case["questions"]) == 3
    module = importlib.import_module("case_contract")
    module.validate_case(case)
    importer = importlib.import_module("import_cases")
    imported = tmp_path / "imported"
    report = importer.import_export(output / "capture_export.json", imported)
    assert report["cases"][0]["effective_split"] == "locked_test"
    assert report["counts"]["train_candidates"] == 0
    assert json.loads((imported / "data/cases/FORM-SYN-001.json").read_text(encoding="utf-8")) == case
    result = subprocess.run([sys.executable, str(PACK / "run_brief.py"),
        str(output / "case_record.json"), "--dry-run"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert "POST-EVENT" not in str(json.loads(result.stdout))
    brief = importlib.import_module("run_brief")
    scorer = importlib.import_module("score_brief")
    response = json.loads((PACK / "tests/fixtures/SYN-001.mock_brief.json").read_text(encoding="utf-8"))
    card = scorer.render_scorecard(case, brief.create_run(case, mock_response=response))
    assert "Invents a measurement\nRecommends an unverified operating change" in card


def test_parser_rejects_duplicate_case_sections_and_unknown_layout(tmp_path):
    module = forms()
    doc = Document(FIXTURE)
    identity = next(t for t in doc.tables if len(t.rows) > 1 and t.cell(1, 0).text == "Case ID (Brad assigns)")
    doc.element.body.insert(-1, copy.deepcopy(identity._element))
    path = tmp_path / "duplicate.docx"
    doc.save(path)
    with pytest.raises(ValueError): module.parse_form(path)
    doc = Document()
    doc.add_paragraph("Unrelated document")
    path = tmp_path / "unrelated.docx"
    doc.save(path)
    with pytest.raises(ValueError): module.parse_form(path)


def test_reviewed_case_id_must_match_any_id_already_on_form():
    with pytest.raises(ValueError, match="case_id.*match"):
        forms().parse_form(FIXTURE, case_id="DIFFERENT", family_id="F-001")


@pytest.mark.parametrize("signature", ["", "unknown", "Illegible signature", "Reviewer / 2026-02-30"])
def test_unparseable_signature_needs_explicit_name_and_date_to_sign(tmp_path, signature):
    doc = Document(FIXTURE)
    row = next(row for table in doc.tables for row in table.rows if row.cells[0].text == "Reviewer signature / date")
    row.cells[1].text = signature
    path = tmp_path / "signature.docx"
    doc.save(path)
    case, _, review = parse_path(path)
    assert case["reviewer_signoff"] == {"signed": False, "name": "", "date": ""}
    assert "reviewer_signoff.name" in review["empty_boxes"]
    with pytest.raises(ValueError, match="reviewer|empty"):
        parse_path(path, case_signed_off=True)
    signed, _, _ = parse_path(path, case_signed_off=True,
                              reviewer_name="Explicit synthetic reviewer", reviewer_date="2026-09-24")
    assert signed["reviewer_signoff"] == {"signed": True, "name": "Explicit synthetic reviewer", "date": "2026-09-24"}


@pytest.mark.parametrize("kwargs", [{"reviewer_name": "Named reviewer"}, {"reviewer_date": "2026-09-24"},
    {"reviewer_name": "unknown", "reviewer_date": "2026-09-24"},
    {"reviewer_name": "Named reviewer", "reviewer_date": "09/24/2026"}])
def test_explicit_signoff_requires_complete_known_name_and_iso_date(kwargs):
    with pytest.raises(ValueError, match="reviewer"):
        parse(**kwargs)


def test_blank_original_consent_is_reported_without_inventing_a_signature(tmp_path):
    doc = Document(FIXTURE)
    for table in doc.tables:
        for row in table.rows:
            if row.cells[0].text in ("Reviewer signature / date", "Interviewer signature / date"):
                row.cells[1].text = ""
    path = tmp_path / "blank-consent.docx"
    doc.save(path)
    case, _, review = parse_path(path, reviewer_name="Explicit reviewer", reviewer_date="2026-09-24")
    assert case["status"] == "draft"
    assert case["reviewer_signoff"]["signed"] is False
    assert review["consent"] == {"reviewer_signature": "", "interviewer_signature": ""}
    assert "consent.reviewer_signature" in review["empty_boxes"]
    assert "consent.interviewer_signature" in review["empty_boxes"]


def test_cli_requires_explicit_complete_family_source_and_publishes_nothing(tmp_path):
    output = tmp_path / "parsed"
    result = subprocess.run([sys.executable, str(PACK / "scripts/case_from_form_fallback.py"), str(FIXTURE),
        str(output), "--case-id", "FORM-SYN-001", "--family-id", "FORM-FAMILY-001"], capture_output=True, text=True)
    assert result.returncode == 2
    assert "--base-export" in result.stderr and "--standalone-family" in result.stderr
    assert not output.exists()


def base_export(case):
    return {"exported_at": "2026-09-24T00:00:00Z", "workflow": {
        "schema": "broadbridge.workflow/1", "updated_at": "", "answers": {}}, "cases": [case]}


def test_base_export_preserves_family_holdout_across_page_and_fallback(tmp_path):
    # The new fallback asks for train, but its page sibling already belongs in locked test.
    page = json.loads((PACK / "tests/fixtures/SYN-TRAIN-001.json").read_text(encoding="utf-8"))
    page["family_id"] = "FORM-FAMILY-001"
    page["questions"][0]["split"] = "locked_test"
    case, _, review = parse(case_signed_off=True)
    for question in case["questions"]:
        question["split"] = "train"
    export = forms().build_export(case, review, base=base_export(page))
    assert export["cases"] == [page, case]
    importlib.import_module("case_contract").validate_export(export)
    path = tmp_path / "combined.json"
    path.write_text(json.dumps(export), encoding="utf-8")
    output = tmp_path / "imported"
    report = importlib.import_module("import_cases").import_export(path, output)
    assert report["counts"]["train_candidates"] == 0
    assert [row["effective_split"] for row in report["cases"]] == ["locked_test", "locked_test"]


@pytest.mark.parametrize("damage", ["duplicate", "invalid_member", "workflow_conflict"])
def test_base_export_conflicts_refuse_instead_of_dropping_cases_or_answers(damage):
    case, _, review = parse()
    page = json.loads((PACK / "tests/fixtures/SYN-TRAIN-001.json").read_text(encoding="utf-8"))
    base = base_export(page)
    if damage == "duplicate": page["case_id"] = case["case_id"].lower()
    if damage == "invalid_member": page["questions"][0]["type"] = "unknown"
    if damage == "workflow_conflict": base["workflow"]["answers"]["A1"] = "Different reviewed answer"
    with pytest.raises(ValueError):
        forms().build_export(case, review, base=base)


def test_fallback_requires_full_family_attestation_for_standalone_export():
    case, _, review = parse()
    with pytest.raises(ValueError, match="family"):
        forms().build_export(case, review)
    export = forms().build_export(case, review, standalone_family=True)
    assert export["cases"] == [case]
    assert export["workflow"]["answers"] == review["workflow"]


@pytest.mark.parametrize("consumer", ["brief", "scorecard", "importer"])
def test_downstream_consumers_refuse_retired_schema(tmp_path, consumer):
    case, _, _ = parse()
    case["schema"] = "broadbridge.case_form/1"
    with pytest.raises(ValueError, match="broadbridge.case_record/1"):
        if consumer == "brief":
            importlib.import_module("run_brief").run_case(case, dry_run=True)
        elif consumer == "scorecard":
            importlib.import_module("score_brief").render_scorecard(case, {})
        else:
            path = tmp_path / "legacy.json"
            path.write_text(json.dumps(base_export(case)), encoding="utf-8")
            importlib.import_module("import_cases").import_export(path, tmp_path / "output")
