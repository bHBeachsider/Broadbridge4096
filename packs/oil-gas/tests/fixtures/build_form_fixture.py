"""Make a clearly synthetic filled copy of the supplied v1 guide; never edit it."""
from pathlib import Path
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[4]
SOURCE = ROOT / "output/expert-capture/Broadbridge_Expert_Interview_Guide_v1.docx"
DEST = Path(__file__).with_name("Synthetic_Filled_Interview_Guide.docx")
document = Document(SOURCE)
document.paragraphs[0].text = "SYNTHETIC TEST ONLY — NO REAL CASE OR CONSENT"
answers = {
    "A1": "Synthetic workflow: collect facts, check instruments, compare mechanisms.",
    "B1": "Synthetic differential pressure increased; establish which checks are needed.",
    "B2": "Synthetic steady feed; no recent instrument work.",
    "B3": "Request wash-oil flow trend and differential-pressure calibration records.",
    "B5": "Possible instrument error or restricted packing; neither is established.",
    "B6": "unknown",
    "B8": "POST-EVENT turning-point inspection result.",
    "B9": "POST-EVENT calculation with independently checked units.",
    "B10": "POST-EVENT historical actions; no recommendation.",
    "B11": "POST-EVENT confidence was limited by unavailable inspection records.",
    "B12": "POST-EVENT dangerous wrong answer would assert certainty.",
    "B13": "POST-EVENT lesson limited to this synthetic example.",
}
identity = ["FORM-SYN-001", "Synthetic form parsing case", "Synthetic vacuum wash section",
            "Synthetic period", "hypothetical", "Synthetic fixture; no customer data", "training"]
questions = [
    ["Missing data", "Which measurements are needed next?", "B1, B4", "SCORING ANSWER ONE: request verified flow and instrument checks.",
     "unknown", "Invents a measurement\nRecommends an unverified operating change", "train"],
    ["Calculation", "What does the supplied pressure difference establish?", "B4", "SCORING ANSWER TWO: insufficient flow information.",
     "unknown", "Fabricates flow", "dev"],
    ["Abstention", "Can the cause be established from the supplied evidence?", "B1, B4", "SCORING ANSWER THREE: cannot establish cause.",
     "unknown", "Claims a proven cause", "locked test"],
]
active = None
for element in document.element.body:
    if element.tag.endswith("}p"):
        text = Paragraph(element, document).text.strip()
        field = re.match(r"([AB]\d{1,2})\s+", text)
        question = re.fullmatch(r"Question (\d+)", text)
        if field: active = field.group(1)
        elif question: active = "Q" + question.group(1)
        elif text == "Identity": active = "identity"
        elif text == "Consent": active = "consent"
    elif element.tag.endswith("}tbl") and active:
        table = Table(element, document)
        if active == "identity":
            for row, value in zip(table.rows[1:], identity): row.cells[1].text = value
        elif active == "consent":
            table.cell(1, 1).text = "SYNTHETIC REVIEWER ONLY / 2026-09-24"
            table.cell(2, 1).text = "SYNTHETIC INTERVIEWER ONLY / 2026-09-24"
        elif active.startswith("Q"):
            for row, value in zip(table.rows[1:], questions[int(active[1:])-1]):
                cell = row.cells[1]
                cell.text = ""
                for i, paragraph in enumerate(value.split("\n")):
                    if i == 0: cell.paragraphs[0].text = paragraph
                    else: cell.add_paragraph(paragraph)
        elif active == "B4":
            for cell, value in zip(table.rows[1].cells, ["T-0", "Bed differential pressure", "18 kPa differential", "Synthetic PDT-101", "unknown"]): cell.text = value
        elif active == "B7":
            for cell, value in zip(table.rows[1].cells, ["POST-EVENT synthetic hypothesis", "POST-EVENT supporting result", "POST-EVENT conflicting result", "POST-EVENT discriminating inspection"]): cell.text = value
        elif active == "B14":
            for i, row in enumerate(table.rows[1:]):
                for j in range(1, 4): row.cells[j].text = "unknown"
                if i == 0:
                    row.cells[1].text = "Synthetic drawing"
                    row.cells[2].text = "yes"
                    row.cells[3].text = "Synthetic fixture only"
        else:
            table.cell(0, 0).text = answers.get(active, "unknown")
        active = None
document.save(DEST)
print(DEST)
