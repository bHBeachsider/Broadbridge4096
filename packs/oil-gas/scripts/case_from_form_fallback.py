"""Fallback DOCX intake into the canonical Case Capture record and complete export."""
import argparse
import copy
from datetime import date, datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sys
import tempfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from case_contract import read_json, validate_case, validate_export
from import_cases import _validate_ids

B_FIELDS = {1: "b1_trigger", 2: "b2_operating_context", 3: "b3_initial_info_and_requests",
            4: "observations", 5: "b5_initial_hypotheses", 6: "b6_distrusted_or_missing",
            7: "hypotheses", 8: "b8_turning_point", 9: "b9_calculations", 10: "b10_actions_taken",
            11: "b11_confidence", 12: "b12_dangerous_wrong_answer", 13: "b13_lesson_and_limits", 14: "evidence"}
IDENTITY = ["Case ID (Brad assigns)", "Short title", "Unit / service (generic label OK)",
            "Approximate date or period", "Real event / reconstructed / hypothetical",
            "Confidentiality and who may see it", "May this case be used for: training / testing only / reference only"]
QUESTION = ["Type (from table above)", "Question as the engineer would ask it", "Evidence supplied with the question (IDs)",
            "Reference answer (Bill)", "Numeric tolerance / units, if any", "Hard-fail: an answer is wrong if it…",
            "Split: train / dev / locked test"]
ROW_FIELDS = {4: ["time", "variable_location", "value_units_basis", "source", "quality"],
              7: ["hypothesis", "evidence_for", "evidence_against", "discriminator"],
              14: ["item", "format", "available_at_decision_time", "restriction"]}
ROW_HEADERS = {4: ["Time", "Variable and location", "Value, units, basis", "Source", "Quality / doubts"],
               7: ["Hypothesis", "Evidence for", "Evidence against", "What would settle it"],
               14: ["Item", "Format", "Available at decision time?", "Sharing restriction"]}


def _normal(text):
    return " ".join(text.split()).casefold()


def _grid(table, labels):
    if len(table.columns) != 2:
        raise ValueError("Unsupported form grid: expected two columns")
    result = {}
    for row in table.rows[1:]:
        label = _normal(row.cells[0].text)
        if label in result:
            raise ValueError(f"Duplicate form field: {label}")
        result[label] = row.cells[1]
    if set(result) != {_normal(label) for label in labels}:
        raise ValueError("Form labels changed or fields are missing; use Interview Guide v1")
    return [result[_normal(label)] for label in labels]


def _extract_tables(document):
    """Anchor by printed labels and question numbers, never document table indexes."""
    found, active, section = {}, None, None
    for element in document.element.body:
        if element.tag.endswith("}p"):
            text = Paragraph(element, document).text.strip()
            part = re.match(r"Part ([ABCD])\s+[—–-]", text)
            marker = re.match(r"([AB])(\d{1,2})\s+", text)
            question = re.fullmatch(r"Question\s+(\d+)", text)
            if part:
                if active: raise ValueError(f"Missing answer table for {active}")
                section, active = part.group(1), None
            elif marker and marker.group(1) == section:
                if active: raise ValueError(f"Missing answer table for {active}")
                active = marker.group(1) + marker.group(2)
            elif text == "Identity" and section == "B":
                active = "identity"
            elif question and section == "C":
                if active: raise ValueError(f"Missing answer table for {active}")
                active = "Q" + question.group(1)
            elif text == "Consent" and section == "D":
                active = "consent"
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            if active:
                if active in found:
                    raise ValueError(f"Repeated {active}; use one case per DOCX")
                found[active], active = table, None
            elif section == "C" and [_normal(c.text) for c in table.rows[0].cells] == ["type", "what it tests", "example shape"]:
                continue  # Printed examples, never answers.
            else:
                raise ValueError("Unmapped table or repeated case; use one case per DOCX")
    required = {"identity", "consent", *(f"A{i}" for i in range(1, 11)), *(f"B{i}" for i in range(1, 15))}
    if active or not required <= found.keys() or not any(key.startswith("Q") for key in found):
        raise ValueError(f"Incomplete Interview Guide v1 layout: {sorted(required - found.keys())}")
    return found


def _known(value):
    return isinstance(value, str) and value.strip().casefold() not in ("", "unknown", "undecided")


def _reviewer(signature, name, signed_date):
    def valid_date(value):
        try:
            return bool(re.fullmatch(r"\d{4}-\d{2}-\d{2}", value)) and date.fromisoformat(value).isoformat() == value
        except ValueError:
            return False

    if name is not None or signed_date is not None:
        if not _known(name) or not isinstance(signed_date, str) or not valid_date(signed_date):
            raise ValueError("Explicit reviewer mapping requires both a known reviewer name and a YYYY-MM-DD reviewer date")
        return name.strip(), signed_date, "explicit_arguments"
    match = re.fullmatch(r"(.+?)\s*/\s*(\d{4}-\d{2}-\d{2})", signature)
    if match and _known(match[1]) and valid_date(match[2]):
        return match[1].strip(), match[2], "docx_signature_date"
    return "", "", "unresolved"


def parse_form(path, *, case_id=None, family_id=None, case_signed_off=False,
               reviewer_name=None, reviewer_date=None):
    path = Path(path)
    tables = _extract_tables(Document(path))
    empty, unused = [], []
    if not _known(case_id) or not _known(family_id):
        raise ValueError("Explicit reviewed case_id and family_id are required before canonical publication")
    if case_id != case_id.strip() or family_id != family_id.strip():
        raise ValueError("Reviewed case_id and family_id cannot contain surrounding whitespace")

    def answer(text, field):
        text = text.strip()
        if not text:
            empty.append(field)
        return text

    def enum(text, field, allowed):
        value = answer(text, field)
        value = value.lower().replace(" ", "_")
        if value not in allowed:
            raise ValueError(f"{field}: explicit choice required from {', '.join(sorted(allowed))}; received {text!r}")
        return value

    workflow = {}
    for i in range(1, 11):
        table = tables[f"A{i}"]
        if len(table.rows) != 1 or len(table.columns) != 1:
            raise ValueError(f"A{i}: expected one answer box")
        workflow[f"A{i}"] = answer(table.cell(0, 0).text, f"workflow.A{i}")
    cells = _grid(tables["identity"], IDENTITY)
    captured_id = cells[0].text.strip()
    if captured_id and captured_id != case_id:
        raise ValueError("Reviewed case_id must match the case ID already written on the form")
    identity = {key: answer(cells[i].text, f"identity.{key}") for i, key in
                [(1, "title"), (2, "unit_service"), (3, "period"), (5, "confidentiality")]}
    identity["record_type"] = enum(cells[4].text, "identity.record_type", {"real_event", "reconstructed", "hypothetical"})
    identity["permitted_use"] = enum(cells[6].text, "identity.permitted_use", {"training", "testing_only", "reference_only"})
    decision, hindsight, evidence = {}, {}, []
    for number, key in B_FIELDS.items():
        table = tables[f"B{number}"]
        prefix = "decision_time." if number <= 6 else "hindsight." if number <= 13 else ""
        field = prefix + key
        if number in ROW_FIELDS:
            if [_normal(c.text) for c in table.rows[0].cells] != [_normal(v) for v in ROW_HEADERS[number]]:
                raise ValueError(f"B{number}: observation/evidence columns changed")
            value = []
            for index, row in enumerate(table.rows[1:]):
                if all(not c.text.strip() for c in row.cells):
                    unused.append(f"{field}[form_row={index+1}]")
                    continue
                value.append({name: answer(cell.text, f"{field}[{len(value)}].{name}")
                              for name, cell in zip(ROW_FIELDS[number], row.cells)})
            if not value: empty.append(field)
        else:
            if len(table.rows) != 1 or len(table.columns) != 1:
                raise ValueError(f"B{number}: expected one answer box")
            value = answer(table.cell(0, 0).text, field)
        if number <= 6: decision[key] = value
        elif number <= 13: hindsight[key] = value
        else: evidence = value
    questions = []
    for key in sorted((k for k in tables if k.startswith("Q")), key=lambda k: int(k[1:])):
        cells = _grid(tables[key], QUESTION)
        if all(not c.text.strip() for c in cells):
            unused.append(f"questions.{key}")
            continue
        prefix = f"questions[{len(questions)}]"
        questions.append({"question_id": f"{case_id}-{key}",
            "type": enum(cells[0].text, prefix+".type", {"brief", "missing_data", "calculation", "grounded_explanation", "abstention"}),
            "question": answer(cells[1].text, prefix+".question"),
            "evidence_ids": answer(cells[2].text, prefix+".evidence_ids"),
            "reference_answer": answer(cells[3].text, prefix+".reference_answer"),
            "tolerance": answer(cells[4].text, prefix+".tolerance"),
            "hard_fail_criteria": answer(cells[5].text, prefix+".hard_fail_criteria"),
            "split": enum(cells[6].text, prefix+".split", {"train", "dev", "locked_test"})})
    if not questions: empty.append("questions")
    consent_table = tables["consent"]
    signature_labels = ["Reviewer signature / date", "Interviewer signature / date"]
    signature_cells = _grid(consent_table, signature_labels)
    consent = {key: cell.text.strip() for key, cell in
               zip(("reviewer_signature", "interviewer_signature"), signature_cells)}
    name, signed_date, signoff_source = _reviewer(consent["reviewer_signature"], reviewer_name, reviewer_date)
    if not name: empty.append("reviewer_signoff.name")
    if not signed_date: empty.append("reviewer_signoff.date")
    if case_signed_off:
        if empty: raise ValueError(f"Cannot mark signed: empty boxes {', '.join(empty)}")
    # Keep original consent gaps visible even when explicit arguments supply the
    # canonical reviewer name/date; the original signatures remain untouched.
    empty.extend("consent." + key for key, value in consent.items() if not value)
    timestamp = datetime.now(timezone.utc).isoformat()
    record = {"schema": "broadbridge.case_record/1", "case_id": case_id, "family_id": family_id,
              "status": "signed" if case_signed_off else "draft", "identity": identity,
              "decision_time": decision, "hindsight": hindsight, "evidence": evidence,
              "questions": questions, "reviewer_signoff": {"signed": case_signed_off, "name": name, "date": signed_date},
              "created_at": timestamp, "updated_at": timestamp}
    validate_case(record)
    _validate_ids([record])
    report = {"workflow": workflow, "consent": consent, "status": record["status"],
              "empty_boxes": empty, "unused_rows_or_questions": unused,
              "signoff_mapping": signoff_source, "case_approval_attested": case_signed_off,
              "timestamp_basis": "canonical fallback capture time; not the historical event or signature date",
              "source": {"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}}
    return record, questions, report


def build_export(case, report, *, base=None, standalone_family=False):
    """Add fallback capture only with a complete export or a separate-family attestation."""
    if (base is None) == (not standalone_family):
        raise ValueError("Supply the complete --base-export or attest --standalone-family before exporting this family")
    validate_case(case)
    if base is None:
        cases, answers = [], {}
    else:
        validate_export(base)  # Never drop invalid family members to make a partial export.
        cases, answers = copy.deepcopy(base["cases"]), dict(base["workflow"]["answers"])
    for key, value in report["workflow"].items():
        previous = answers.get(key, "")
        if previous.strip() and value.strip() and previous != value:
            raise ValueError(f"Workflow conflict in {key}; review the complete export before combining captures")
        if value.strip() or key not in answers:
            answers[key] = value
    cases.append(copy.deepcopy(case))
    _validate_ids(cases)
    timestamp = datetime.now(timezone.utc).isoformat()
    export = {"exported_at": timestamp, "workflow": {"schema": "broadbridge.workflow/1",
              "answers": answers, "updated_at": timestamp}, "cases": cases}
    validate_export(export)
    return export


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx")
    parser.add_argument("out_dir")
    parser.add_argument("--case-id", required=True, help="Explicitly reviewed case ID; must match a populated guide ID")
    parser.add_argument("--family-id", required=True, help="Reviewed case-family ID; the guide has no family box")
    parser.add_argument("--case-signed-off", action="store_true", help="Attest Bill approved this written case; Part D consent alone is insufficient")
    parser.add_argument("--reviewer-name", help="Explicit reviewer name if the DOCX signature cannot be parsed")
    parser.add_argument("--reviewer-date", help="Explicit YYYY-MM-DD signoff date; supply with --reviewer-name")
    family = parser.add_mutually_exclusive_group(required=True)
    family.add_argument("--base-export", help="Complete all-records export including every existing related case; preserve all members")
    family.add_argument("--standalone-family", action="store_true", help="Attest this case is the entire family and has no related cases outside this capture")
    args = parser.parse_args(argv)
    try:
        case, questions, report = parse_form(args.docx, case_id=args.case_id, family_id=args.family_id,
            case_signed_off=args.case_signed_off, reviewer_name=args.reviewer_name, reviewer_date=args.reviewer_date)
        base = read_json(args.base_export) if args.base_export else None
        export = build_export(case, report, base=base, standalone_family=args.standalone_family)
        report["family_scope"] = {"basis": "complete_base_export" if args.base_export else "standalone_family_attestation",
            "base_export": str(Path(args.base_export).resolve()) if args.base_export else None,
            "case_count": len(export["cases"])}
        output = Path(args.out_dir).absolute()
        if output.exists() or output.is_symlink():
            raise ValueError("out_dir must not exist; preserve previous review snapshots")
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".form-import-", dir=output.parent) as temp:
            stage = Path(temp) / "snapshot"
            stage.mkdir()
            for filename, value in [("case_record.json", case), ("capture_export.json", export), ("form_review.json", report)]:
                (stage / filename).write_text(json.dumps(value, ensure_ascii=False, indent=2)+"\n", encoding="utf-8", newline="\n")
            stage.rename(output)
        print(json.dumps({"case_id": case["case_id"], "status": case["status"], "questions": len(questions),
                          "empty_boxes": report["empty_boxes"], "unused_rows_or_questions": report["unused_rows_or_questions"]}))
        return 0
    except (ValueError, OSError) as exc:
        print(f"FORM REFUSED: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
