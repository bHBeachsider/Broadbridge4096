"""Parse one filled Interview Guide v1 into a case and Part C questions."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import sys
import tempfile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from form_contract import validate_form_case

B_FIELDS = {1: "b1_trigger", 2: "b2_operating_context", 3: "b3_initial_info_and_requests",
            4: "b4_observations", 5: "b5_initial_hypotheses", 6: "b6_distrusted_or_missing",
            7: "b7_hypotheses", 8: "b8_turning_point", 9: "b9_calculations", 10: "b10_actions_taken",
            11: "b11_confidence", 12: "b12_dangerous_wrong_answer", 13: "b13_lesson_and_limits", 14: "b14_evidence"}
IDENTITY = ["Case ID (Brad assigns)", "Short title", "Unit / service (generic label OK)",
            "Approximate date or period", "Real event / reconstructed / hypothetical",
            "Confidentiality and who may see it", "May this case be used for: training / testing only / reference only"]
QUESTION = ["Type (from table above)", "Question as the engineer would ask it", "Evidence supplied with the question (IDs)",
            "Reference answer (Bill)", "Numeric tolerance / units, if any", "Hard-fail: an answer is wrong if it…",
            "Split: train / dev / locked test"]
ROW_FIELDS = {4: ["time", "variable_location", "value_units_basis", "source", "quality"],
              7: ["hypothesis", "evidence_for", "evidence_against", "discriminator"],
              14: ["item", "format", "available_at_decision_time", "restriction"]}
ROW_HEADERS = {4: ["Time", "Variable and location", "Value, units, basis", "Source", "Quality / doubts"],
               7: ["Hypothesis", "Evidence for", "Evidence against", "What would settle it"],
               14: ["Item", "Format", "Available at decision time?", "Sharing restriction"]}


def _normal(text):
    return " ".join(text.split()).casefold()


def _grid(table, labels):
    if len(table.columns) != 2:
        raise ValueError("Unsupported form grid: expected two columns")
    result = {}
    for row in table.rows[1:]:
        label = _normal(row.cells[0].text)
        if label in result:
            raise ValueError(f"Duplicate form field: {label}")
        result[label] = row.cells[1]
    if set(result) != {_normal(label) for label in labels}:
        raise ValueError("Form labels changed or fields are missing; use Interview Guide v1")
    return [result[_normal(label)] for label in labels]


def _extract_tables(document):
    """Anchor by printed labels and question numbers, never document table indexes."""
    found, active, section = {}, None, None
    for element in document.element.body:
        if element.tag.endswith("}p"):
            text = Paragraph(element, document).text.strip()
            part = re.match(r"Part ([ABCD])\s+[—–-]", text)
            marker = re.match(r"([AB])(\d{1,2})\s+", text)
            question = re.fullmatch(r"Question\s+(\d+)", text)
            if part:
                if active: raise ValueError(f"Missing answer table for {active}")
                section, active = part.group(1), None
            elif marker and marker.group(1) == section:
                if active: raise ValueError(f"Missing answer table for {active}")
                active = marker.group(1) + marker.group(2)
            elif text == "Identity" and section == "B":
                active = "identity"
            elif question and section == "C":
                if active: raise ValueError(f"Missing answer table for {active}")
                active = "Q" + question.group(1)
            elif text == "Consent" and section == "D":
                active = "consent"
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            if active:
                if active in found:
                    raise ValueError(f"Repeated {active}; use one case per DOCX")
                found[active], active = table, None
            elif section == "C" and [_normal(c.text) for c in table.rows[0].cells] == ["type", "what it tests", "example shape"]:
                continue  # Printed examples, never answers.
            else:
                raise ValueError("Unmapped table or repeated case; use one case per DOCX")
    required = {"identity", "consent", *(f"A{i}" for i in range(1, 11)), *(f"B{i}" for i in range(1, 15))}
    if active or not required <= found.keys() or not any(key.startswith("Q") for key in found):
        raise ValueError(f"Incomplete Interview Guide v1 layout: {sorted(required - found.keys())}")
    return found


def parse_form(path, *, family_id=None, case_signed_off=False):
    path = Path(path)
    tables = _extract_tables(Document(path))
    empty, unused = [], []

    def answer(text, field):
        text = text.strip()
        if not text:
            empty.append(field)
            return None
        return "unknown" if text.casefold() == "unknown" else text

    def enum(text, field, allowed):
        value = answer(text, field)
        if value is None: return None
        value = value.lower().replace(" ", "_")
        if value not in {*allowed, "unknown"}:
            raise ValueError(f"{field}: unrecognized choice {text!r}")
        return value

    workflow = {}
    for i in range(1, 11):
        table = tables[f"A{i}"]
        if len(table.rows) != 1 or len(table.columns) != 1:
            raise ValueError(f"A{i}: expected one answer box")
        workflow[f"A{i}"] = answer(table.cell(0, 0).text, f"workflow.A{i}")
    cells = _grid(tables["identity"], IDENTITY)
    case_id = answer(cells[0].text, "case_id")
    family_id = answer(family_id or "", "family_id")
    identity = {key: answer(cells[i].text, f"identity.{key}") for i, key in
                [(1, "title"), (2, "unit_service"), (3, "period"), (5, "confidentiality")]}
    identity["record_type"] = enum(cells[4].text, "identity.record_type", {"real_event", "reconstructed", "hypothetical"})
    identity["permitted_use"] = enum(cells[6].text, "identity.permitted_use", {"training", "testing_only", "reference_only", "undecided"})
    decision, hindsight, evidence = {}, {}, []
    for number, key in B_FIELDS.items():
        table = tables[f"B{number}"]
        prefix = "decision_time." if number <= 6 else "hindsight." if number <= 13 else ""
        field = prefix + key
        if number in ROW_FIELDS:
            if [_normal(c.text) for c in table.rows[0].cells] != [_normal(v) for v in ROW_HEADERS[number]]:
                raise ValueError(f"B{number}: observation/evidence columns changed")
            value = []
            for index, row in enumerate(table.rows[1:]):
                if all(not c.text.strip() for c in row.cells):
                    unused.append(f"{field}[form_row={index+1}]")
                    continue
                value.append({name: answer(cell.text, f"{field}[{len(value)}].{name}")
                              for name, cell in zip(ROW_FIELDS[number], row.cells)})
            if not value: empty.append(field)
        else:
            if len(table.rows) != 1 or len(table.columns) != 1:
                raise ValueError(f"B{number}: expected one answer box")
            value = answer(table.cell(0, 0).text, field)
        if number <= 6: decision[key] = value
        elif number <= 13: hindsight[key] = value
        else: evidence = value
    questions = []
    for key in sorted((k for k in tables if k.startswith("Q")), key=lambda k: int(k[1:])):
        cells = _grid(tables[key], QUESTION)
        if all(not c.text.strip() for c in cells):
            unused.append(f"questions.{key}")
            continue
        prefix = f"questions[{len(questions)}]"
        supplied_ids = answer(cells[2].text, prefix + ".evidence_ids")
        criteria = [p.text.strip() for p in cells[5].paragraphs if p.text.strip()]
        if not criteria: empty.append(prefix + ".hard_fail_criteria")
        questions.append({"case_id": case_id, "family_id": family_id,
            "question_id": f"{case_id}-{key}" if case_id not in (None, "unknown") else None,
            "type": enum(cells[0].text, prefix+".type", {"brief", "missing_data", "calculation", "grounded_explanation", "abstention"}),
            "question": answer(cells[1].text, prefix+".question"),
            "evidence_ids": [v.strip() for v in re.split(r"[,\n]", supplied_ids) if v.strip()] if supplied_ids else [],
            "reference_answer": answer(cells[3].text, prefix+".reference_answer"),
            "tolerance": answer(cells[4].text, prefix+".tolerance"),
            "hard_fail_criteria": ["unknown" if v.casefold() == "unknown" else v for v in criteria],
            "split": enum(cells[6].text, prefix+".split", {"train", "dev", "locked_test"})})
    if not questions: empty.append("questions")
    consent_table = tables["consent"]
    signature_labels = ["Reviewer signature / date", "Interviewer signature / date"]
    signature_cells = _grid(consent_table, signature_labels)
    consent = {key: answer(cell.text, "consent."+key) for key, cell in
               zip(("reviewer_signature", "interviewer_signature"), signature_cells)}
    if case_signed_off:
        if empty: raise ValueError(f"Cannot mark signed: empty boxes {', '.join(empty)}")
        required_values = [case_id, family_id, consent["reviewer_signature"], identity["permitted_use"]]
        if any(value in (None, "unknown", "undecided") for value in required_values):
            raise ValueError("Cannot mark signed: identity, family, reviewer or permission remains unknown")
    record = {"schema": "broadbridge.case_form/1", "case_id": case_id, "family_id": family_id,
              "status": "signed" if case_signed_off else "pending_review", "identity": identity,
              "decision_time": decision, "hindsight": hindsight, "b14_evidence": evidence,
              "questions": questions, "reviewer_signoff": {"signed": case_signed_off, **consent,
                  "approval_basis": "operator_attestation" if case_signed_off else "pending_review"},
              "source": {"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}}
    validate_form_case(record)
    report = {"workflow": workflow, "consent": consent, "status": record["status"],
              "empty_boxes": empty, "unused_rows_or_questions": unused, "source": record["source"]}
    return record, questions, report


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("docx")
    parser.add_argument("out_dir")
    parser.add_argument("--family-id", help="Reviewed case-family ID; the guide has no family box")
    parser.add_argument("--case-signed-off", action="store_true", help="Attest Bill approved this written case; Part D consent alone is insufficient")
    args = parser.parse_args(argv)
    try:
        case, questions, report = parse_form(args.docx, family_id=args.family_id, case_signed_off=args.case_signed_off)
        output = Path(args.out_dir).absolute()
        if output.exists() or output.is_symlink():
            raise ValueError("out_dir must not exist; preserve previous review snapshots")
        output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix=".form-import-", dir=output.parent) as temp:
            stage = Path(temp) / "snapshot"
            stage.mkdir()
            for filename, value in [("case_record.json", case), ("eval_questions.json", questions), ("form_review.json", report)]:
                (stage / filename).write_text(json.dumps(value, ensure_ascii=False, indent=2)+"\n", encoding="utf-8", newline="\n")
            stage.rename(output)
        print(json.dumps({"case_id": case["case_id"], "status": case["status"], "questions": len(questions),
                          "empty_boxes": report["empty_boxes"], "unused_rows_or_questions": report["unused_rows_or_questions"]}))
        return 0
    except (ValueError, OSError) as exc:
        print(f"FORM REFUSED: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
