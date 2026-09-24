from pathlib import Path
import csv
import json
import re
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont
from build_oil_gas_documents import base, table as word_table

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / 'scripts/multimodal_cto_plan_content.json').read_text(encoding='utf-8'))
REFS = DATA['refs']
OUT = ROOT / 'output/docx/Broadbridge_Oil_and_Gas_Multimodal_SLM_CTO_Plan.docx'
REG = ROOT / 'output/multimodal-project'
QA = ROOT / 'tmp/docx-qa/multimodal-cto'
REG.mkdir(parents=True, exist_ok=True)
QA.mkdir(parents=True, exist_ok=True)

def diagram():
    im = Image.new('RGB', (1800, 760), 'white')
    dr = ImageDraw.Draw(im)
    fp = Path('C:/Windows/Fonts/arial.ttf')
    bold = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 28)
    font = ImageFont.truetype(str(fp), 25)
    small = ImageFont.truetype(str(fp), 23)
    ink, border, fill = '#18364D', '#78909E', '#EDF3F6'
    def box(x, y, w, h, title, lines):
        dr.rounded_rectangle((x,y,x+w,y+h), radius=10, fill=fill, outline=border, width=3)
        dr.text((x+18,y+14), title, font=bold, fill=ink)
        for j, line in enumerate(lines):
            dr.text((x+18,y+57+j*31), line, font=font, fill='#222222')
    def arrow(points):
        dr.line(points, fill=ink, width=4)
        x,y=points[-1]; px,py=points[-2]
        if x>px: tri=[(x,y),(x-15,y-9),(x-15,y+9)]
        elif x<px: tri=[(x,y),(x+15,y-9),(x+15,y+9)]
        elif y>py: tri=[(x,y),(x-9,y-15),(x+9,y-15)]
        else: tri=[(x,y),(x-9,y+15),(x+9,y+15)]
        dr.polygon(tri,fill=ink)
    box(30,35,470,145,'Approved source intake',['Expert cases, documents, images','Originals retained in quarantine'])
    box(665,35,470,145,'Admission and extraction',['Rights, quality and malware checks','OCR, regions, units and lineage'])
    box(1300,35,470,145,'Controlled evidence store',['Versioned sources and images','Access groups and source grants'])
    arrow([(500,108),(665,108)])
    arrow([(1135,108),(1300,108)])
    box(30,290,470,165,'Request and retrieval',['Authenticated internal user','Authorized passages and regions','Evidence available at decision time'])
    box(665,290,470,165,'Vision language SLM',['Approved model and adapter','Evidence-linked hypotheses','Allowlisted tool requests'])
    box(1300,290,470,165,'Validated brief and review',['Schema and source checks','Engineer accepts or rejects','Corrections enter review queue'])
    arrow([(1535,180),(1535,230),(265,230),(265,290)])
    arrow([(500,375),(665,375)])
    arrow([(1135,375),(1300,375)])
    box(665,565,470,130,'Checked engineering tools',['Typed inputs, units and ranges','Reproducible method and results'])
    arrow([(845,455),(845,565)])
    arrow([(965,565),(965,455)])
    dr.text((32,578),'Training uses a separate approved snapshot.',font=small,fill='#333333')
    dr.text((32,616),'No control-system write path.',font=small,fill='#333333')
    dr.text((1300,578),'No automatic feedback training.',font=small,fill='#333333')
    dr.text((1300,616),'Technical and security release gates.',font=small,fill='#333333')
    path=QA/'architecture.png'; im.save(path)
    return path

def hyperlink(para,label,url,size=None):
    rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True)
    e=OxmlElement('w:hyperlink'); e.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'18364D'); prop.append(color)
    if size:
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(round(size*2)));prop.append(sz)
    r.append(prop);tx=OxmlElement('w:t');tx.text=label;r.append(tx);e.append(r);para._p.append(e)

def rich(para,text,size=None):
    for part in re.split(r'(\[[A-Z]\d+\])',text):
        key=part[1:-1]
        if part.startswith('[') and key in REFS:
            hyperlink(para,part,REFS[key][1],size)
        else:
            run=para.add_run(part)
            if size:run.font.size=Pt(size)

def markdown(text):
    return re.sub(r'\[([A-Z]\d+)\]',lambda m:'['+m[1]+']('+REFS[m[1]][1]+')',text)

d=base('Multimodal SLM project delivery plan')
d.paragraphs[0].text='Broadbridge Oil and Gas'
d.paragraphs[0].style=d.styles['Title']
d.paragraphs[2].text='CTO program recommendation  |  Broadbridge4096  |  23 September 2026'
d.paragraphs[2].runs[0].font.size=Pt(9)
md=['# Broadbridge Oil and Gas multimodal SLM project delivery plan','CTO program recommendation | 23 September 2026','']
md.append('This plan supersedes earlier text-only and paid-pilot assumptions for the proposed internal multimodal pilot. Financial assumptions and acceptance targets remain proposals.\n')
for idx,page in enumerate(DATA['pages']):
    if idx:d.add_page_break()
    d.add_heading(page['title'],1)
    md.extend(['## '+page['title'],''])
    for block in page['blocks']:
        kind=block['type']
        if kind=='p':
            para=d.add_paragraph();rich(para,block['text'])
            md.extend([markdown(block['text']),''])
        elif kind=='h':
            d.add_heading(block['text'],2);md.extend(['### '+block['text'],''])
        elif kind=='diagram':
            para=d.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pic=para.add_run().add_picture(str(diagram()),width=Inches(7.2))
            pic._inline.docPr.set('descr','Approved source intake passes admission to a controlled evidence store. Authorized retrieval supplies a vision language model, which uses checked tools and produces a validated brief for engineer review.')
            md.extend(['```mermaid','flowchart LR','  S[Approved sources] --> A[Admission and extraction] --> E[Controlled evidence store]','  U[Internal user] --> R[Authorized retrieval]','  E --> R --> M[Vision language SLM] --> B[Validated brief and engineer review]','  M <--> T[Checked engineering tools]','```',''])
        elif kind=='table':
            rows=block['rows']
            if idx==13:
                rows=[[('$'+format(v,',')) if isinstance(v,int) else v for v in row] for row in rows]
            tab=word_table(d,block['headers'],rows,block['widths'],10.2)
            for row in tab.rows[1:]:
                for cell in row.cells:
                    if re.search(r'\[[A-Z]\d+\]',cell.text):
                        txt=cell.text;cell.paragraphs[0].clear();rich(cell.paragraphs[0],txt,10.2)
            md.extend(['| '+' | '.join(block['headers'])+' |','|'+'|'.join(['---']*len(block['headers']))+'|'])
            for row in rows:md.append('| '+' | '.join(markdown(str(v)) for v in row)+' |')
            md.append('')

d.core_properties.title='Broadbridge Oil and Gas Multimodal SLM Project Delivery Plan'
d.core_properties.subject='Internal engineering assistant delivery architecture governance evaluation and budget'
d.core_properties.comments=''
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT)
md.extend(['## Linked reference index',''])
for key,(title,url) in REFS.items():md.extend([f'- [{key} {title}]({url})'])
(ROOT/'Broadbridge-Multimodal-SLM-CTO-Plan.md').write_text('\n'.join(md)+'\n',encoding='utf-8')

for name,headers,rows in [
 ('Broadbridge_Multimodal_Execution_Register.csv',['package_id','work_package','accountable_owner','weeks','dependencies','deliverable','acceptance'],DATA['packages']),
 ('Broadbridge_Multimodal_Gate_Register.csv',['gate','end_week','accountable_owner','required_evidence','required_reviews','decision_rule'],DATA['gates'])]:
    with (REG/name).open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f);writer.writerow(headers);writer.writerows(rows)

assert len(DATA['pages'])==17
assert len(DATA['packages'])==12 and len(DATA['gates'])==6
assert sum([.25,.25,.5,.75,.5,1,1])==4.25
assert sum([240,240,480,720,480,960,960])==4080
assert sum([625,300,225,170,180])==1500
assert sum([275,100,75,30,20])==500
assert sum([408000,200000,30000,20000,30000,20000])*1.2==849600
assert sum([612000,420000,80000,50000,60000,80000])*1.2==1562400
print(json.dumps({'docx':str(OUT),'planned_pages':17,'work_packages':12,'gates':6,'training_examples':2000,'fte':4.25}))
