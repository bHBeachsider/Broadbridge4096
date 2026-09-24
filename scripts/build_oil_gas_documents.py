from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output' / 'docx'
OUT.mkdir(parents=True, exist_ok=True)
STAGES = [
 ('Managing Director',1,1,1), ('Commercial',1,2,3), ('Technical practices',2,4,7),
 ('Product, AI, and knowledge engineering',3,4,6),
 ('Dedicated delivery and customer success',0,2,4), ('Dedicated academy team',0,0,2),
 ('Subsidiary business operations',0,0,1)]
assert [sum(r[i] for r in STAGES) for i in (1,2,3)] == [7,13,24]
ROLES = [
 ('Managing Director','Subsidiary Board','Own strategy, P&L, partnerships, staffing and execution. Approve customer commitments within delegated limits and report operating results to the board.'),
 ('Commercial Director','Managing Director','Own customer development, qualification, proposals, pricing recommendations and renewals. Maintain the pipeline and confirm the buyer, budget and purchasing route.'),
 ('Technical Director','Managing Director','Own engineering quality, technical scope, experts, independent review and technical acceptance. Lead the launch practice and suspend unsupported or unsafe deliverables.'),
 ('Process Applications Engineer','Technical Director','Define customer problems, check engineering inputs, conduct analysis and coordinate pilot delivery. Record assumptions, evidence and customer acceptance requirements.'),
 ('Head of Product & AI','Managing Director','Own product priorities, platform design, model evaluation and releases. Authorize deployment only after required technical and security acceptance.'),
 ('Applied AI Engineer','Head of Product & AI','Build retrieval, engineering tools, integrations and deployment pipelines. Implement access controls, evaluations, monitoring and rollback under product leadership.'),
 ('Knowledge Engineer','Head of Product & AI','Conduct expert interviews, develop cases, maintain source rights records and prepare training content. Submit domain content to the Technical Director for acceptance.')]

def shade(c,fill):
 e=OxmlElement('w:shd'); e.set(qn('w:fill'),fill); c._tc.get_or_add_tcPr().append(e)

def base(subtitle):
 d=Document(); s=d.sections[0]
 s.page_width=Inches(8.5); s.page_height=Inches(11)
 s.top_margin=Inches(.62); s.bottom_margin=Inches(.58)
 s.left_margin=Inches(.65); s.right_margin=Inches(.65)
 s.footer_distance=Inches(.25)
 d.settings.odd_and_even_pages_header_footer=False
 s.different_first_page_header_footer=False
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
  st=d.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor(0,0,0)
  st.font.underline=False
  for border in list(st.element.xpath('./w:pPr/w:pBdr')): border.getparent().remove(border)
 n=d.styles['Normal']; n.font.size=Pt(11); n.paragraph_format.space_after=Pt(6); n.paragraph_format.line_spacing=1.06
 for name,size,before,after in [('Title',25,0,5),('Subtitle',13,0,8),('Heading 1',16,10,7),('Heading 2',12,8,5)]:
  st=d.styles[name]; st.font.size=Pt(size); st.font.bold=name.startswith('Heading')
  st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
 d.add_paragraph('Broadbridge Oil & Gas','Title'); d.add_paragraph(subtitle,'Subtitle')
 p=d.add_paragraph('Broadbridge4096 subsidiary design  |  23 September 2026'); p.runs[0].font.size=Pt(9)
 p=s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 p.add_run('Broadbridge Oil & Gas  |  ').font.size=Pt(9)
 f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); p._p.append(f)
 d.core_properties.author='Broadbridge4096'; d.core_properties.title='Broadbridge Oil & Gas '+subtitle
 return d

def p(d,text): return d.add_paragraph(text)
def h(d,text,level=1): d.add_heading(text,level)
def page(d,title): d.add_page_break(); h(d,title)
def center(d,text,bold=False,size=11):
 para=d.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.space_after=Pt(5)
 r=para.add_run(text); r.bold=bold; r.font.size=Pt(size)

def table(d,headers,rows,widths,size=10.5):
 t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for c,w in zip(t.columns,widths): c.width=Inches(w)
 for row,values in [(t.rows[0],headers)]+[(t.add_row(),v) for v in rows]:
  for c,w,v in zip(row.cells,widths,values):
   c.width=Inches(w); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; c.text=str(v)
   mar=OxmlElement('w:tcMar')
   for edge,val in [('top','80'),('left','100'),('bottom','80'),('right','100')]:
    e=OxmlElement('w:'+edge); e.set(qn('w:w'),val); e.set(qn('w:type'),'dxa'); mar.append(e)
   c._tc.get_or_add_tcPr().append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1.02
    if isinstance(v,int): para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in para.runs: r.font.size=Pt(size)
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for i,row in enumerate(t.rows):
  for c in row.cells:
   shade(c,'18364D' if i==0 else ('F2F4F5' if i%2 else 'FFFFFF'))
   if i==0:
    for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
 t._tbl.tblPr.append(borders)
 para=d.add_paragraph(); para.paragraph_format.space_after=Pt(1); para.paragraph_format.line_spacing=Pt(2); para.add_run().font.size=Pt(2)
 return t

def staffing(d):
 table(d,['Function','Launch','Expansion','Mature target'],STAGES+[('Dedicated employees',7,13,24)],[4.05,1.05,1.05,1.05])

def chart():
 d=base('Subsidiary organizational chart')
 p(d,'Working name. Initial ownership assumption: 100% Broadbridge4096. Incorporation jurisdiction and final legal form remain a subsequent legal implementation decision.')
 h(d,'Launch reporting lines')
 center(d,'Broadbridge4096',True,14)
 center(d,'↓  ownership and reserved decisions',size=10)
 center(d,'Subsidiary Board',True,13)
 center(d,'↓  oversight and management accountability',size=10)
 center(d,'Managing Director  •  1 employee',True,13)
 center(d,'↓  Each leader below reports to the Managing Director',size=10)
 table(d,['Commercial Director\n1 employee','Technical Director\n1 employee','Head of Product & AI\n1 employee'],[
  ('Customer development\nProposals and pricing\nRenewals','Engineering quality\nExperts and review\nLaunch practice leadership','Product priorities\nPlatform and evaluation\nRelease decisions'),
  ('No dedicated direct report at launch','↓ Direct report\nProcess Applications Engineer\n1 employee','↓ Direct reports\nApplied AI Engineer\nKnowledge Engineer\n2 employees')],[2.4,2.4,2.4],11)
 p(d,'Employee reconciliation: 1 Managing Director + 3 functional leaders + 3 individual contributors = 7 dedicated employees. Each employee has one line manager.')
 h(d,'Expert council and independent review',2)
 p(d,'Expert Council (contracted experts) → advisory reporting to Technical Director. The council advises on coverage, quality and limitations and can escalate directly to the Subsidiary Board. It has no management authority. The Technical Director manages expert engagements; independent reviewers cannot certify their own authored work.')
 p(d,'Norman Lieberman is an illustrative founding expert candidate. No appointment, endorsement or content license is implied. Experts and reviewers are contractors, excluded from the employee count.')
 h(d,'Parent shared services',2)
 p(d,'Broadbridge4096 provides finance, treasury, legal, HR, IT security and corporate administration under defined service arrangements. The Managing Director owns the service relationship. Shared-service personnel remain outside subsidiary employee totals.')
 page(d,'Future practices and staffing')
 p(d,'All practice leaders report to the Technical Director. Practices define sector expertise; commercial, product and delivery functions serve them across the subsidiary.')
 table(d,['Practice','Activation and reporting'],[
  ('Downstream and Petrochemicals','Launch: refining, process troubleshooting, equipment performance and operating knowledge. The Technical Director leads initially; a dedicated practice lead is added at expansion.'),
  ('Gas Processing and Midstream','Subsequent: treatment, compression, processing, LNG-related operations and facilities. A practice lead reports to the Technical Director after expansion gates are met.'),
  ('Upstream Production and Facilities','Later: its own specialists and practice lead. Drilling and subsurface services require separately demonstrated capability and scope approval.')],[2.35,4.85])
 p(d,'Each practice can sell advisory engagements, recurring software subscriptions, training programs and licensed expert knowledge within its accepted technical scope.')
 staffing(d)
 h(d,'Future reporting lines',2)
 p(d,'At expansion, the Delivery Lead reports to the Managing Director and manages the Customer Success Manager. At maturity, the Head of Academy and Business Operations Manager also report to the Managing Director. Commercial hires report to the Commercial Director; technical hires to the Technical Director or their practice lead; product, AI and knowledge hires to the Head of Product & AI.')
 p(d,'Launch staff cover delivery and training. Dedicated delivery, academy and operations posts are future positions. The operating structure specifies the exact role mix at each staffing stage.')
 h(d,'Conditions for growth',2)
 p(d,'Expansion requires repeatable paid deployments, funded staffing and a paid customer mandate for every new practice. The mature structure is a destination, not an immediate hiring commitment. Contracted experts, board participation and parent shared services remain outside all employee totals.')
 d.save(OUT/'Broadbridge_Oil_and_Gas_Organizational_Chart.docx')

def operating():
 d=base('Business and operating structure')
 h(d,'1 Mandate and business model')
 p(d,'Broadbridge Oil & Gas will operate as a dedicated subsidiary of Broadbridge4096, initially assumed to be 100% parent-owned. It will combine expert advisory services, AI products, training and knowledge licensing for external oil and gas customers. The commodity-company documents provide organizational examples; the subsidiary has its own customers, product priorities and operating results.')
 p(d,'Broadbridge Oil & Gas is a working name. Incorporation jurisdiction and final legal form will be decided during legal implementation. This design does not depend on either choice and does not represent an incorporation or executed expert appointment.')
 table(d,['Industry practice','Business scope and sequence'],[
  ('Downstream and Petrochemicals','Launch in refining and process operations. Begin with a bounded equipment or troubleshooting domain supported by cleared expert material and paid customer demand.'),
  ('Gas Processing and Midstream','Expand into treatment, compression, processing, LNG-related operations and facilities after securing appropriate specialists and a paid mandate.'),
  ('Upstream Production and Facilities','Enter later with dedicated specialists. Drilling and subsurface services need separately demonstrated capability before being offered.')],[2.3,4.9])
 h(d,'Four offerings within each practice',2)
 table(d,['Offering','Customer purchase and delivery boundary'],[
  ('Advisory engagements','A defined engineering question, scope, assumptions and reviewed deliverable under a statement of work. Complex field work is separately staffed and authorized.'),
  ('Software subscriptions','A validated workflow or module with defined users/sites, support allowance, updates and human escalation. Bespoke integrations are separately priced.'),
  ('Training programs','Instructor-led or digital case-based learning with learning objectives, cleared materials and reviewed exercises. Launch staff and contracted experts deliver the academy offering.'),
  ('Licensed expert knowledge','Specified access or reuse rights to expert cases, methods or content. The license defines the audience, uses, duration and sublicensing boundaries.')],[2,5.2])
 p(d,'Customer value must be demonstrated through observed workflow improvements and accepted technical quality. Published expertise is a starting resource; access, rights, transferability and willingness to pay must be established for each offer.')
 page(d,'2 Launch roles and reporting')
 p(d,'The seven dedicated employees below form the launch team. Contracted experts, independent reviewers, board members and parent shared-service personnel are not additional employee positions.')
 table(d,['Role and line manager','Charter'],[(r+'\nReports to '+m,c) for r,m,c in ROLES],[2.25,4.95])
 h(d,'Delivery and training coverage',2)
 p(d,'At launch, the Technical Director is accountable for customer delivery, with the Process Applications Engineer coordinating each pilot. The Commercial Director owns the customer relationship. The Head of Product & AI owns software implementation, and the Knowledge Engineer prepares training content. The Technical Director approves technical content and instructor capability.')
 p(d,'The Knowledge Engineer has one line manager, the Head of Product & AI. The Technical Director sets domain acceptance requirements and provides technical review; this does not create a second reporting line. Each leader plans capacity across client work and reusable product development.')
 page(d,'3 Governance and decision rights')
 p(d,'Broadbridge4096 appoints the Subsidiary Board. The board oversees strategy, performance and management and escalates reserved matters to the parent. The Managing Director reports to the board. Each decision below has one accountable owner; required reviews are conditions of approval, not shared accountability.')
 table(d,['Decision','Accountable owner','Required input or boundary'],[
  ('Capital, annual budget, equity and material IP transactions','Broadbridge4096','Subsidiary Board recommends; parent approval is required. Materiality is set in the delegation schedule.'),
  ('Senior appointments','Broadbridge4096','Board recommends the Managing Director. The Managing Director recommends functional leaders through the board.'),
  ('Oversight and management performance','Subsidiary Board','Review results, challenge risk and hear expert escalations. Respect parent reserved matters.'),
  ('Customer qualification and pipeline','Commercial Director','Confirm need, budget, procurement and fit. Technical Director confirms capability before a proposal is committed.'),
  ('Pricing, contracts and scope changes','Managing Director','Commercial Director proposes terms. Obtain technical, delivery, legal and security inputs. Commit within delegated limits.'),
  ('Technical scope and acceptance','Technical Director','Obtain independent review and record limitations. May suspend unsupported or unsafe content, advice or releases.'),
  ('Customer delivery','Technical Director at launch; Delivery Lead after appointment','One accountable delivery owner per engagement. Record the handover. Technical acceptance stays with the Technical Director.'),
  ('Product priorities and software release','Head of Product & AI','Technical Director accepts domain behavior; parent IT security accepts security controls. Both are required for release.'),
  ('Security acceptance','Designated parent IT security lead','Assess access, hosting, isolation and incident controls under the shared-service arrangement.'),
  ('Independent review findings','Assigned independent reviewer','Review within their competence. They cannot certify their own source material or implementation.')],[1.8,1.8,3.6],10)
 p(d,'Before commitments begin, the parent approves written limits for contract value, term, nonstandard liability, expenditure and material IP transactions. No monetary limits are assumed here. Until delegation is approved, the Managing Director routes commitments to the parent through the board.')
 p(d,'Commercial deadlines cannot override technical or security stops. The board resolves resourcing and business choices; the relevant acceptance owner must close the underlying issue before release.')
 page(d,'4 Experts and intellectual property')
 p(d,'The Technical Director appoints and manages contracted experts within the approved budget and delegated contracting process. An Expert Council advises on coverage, methods, evidence and limitations and reports to the Technical Director. Members may raise unresolved concerns directly with the Subsidiary Board. Council membership confers advisory influence, not management or signing authority.')
 p(d,'Norman Lieberman is an illustrative founding expert candidate based on his published process-engineering and teaching experience. Participation, availability, content rights and commercial terms are unconfirmed. His name must not imply endorsement or an agreed appointment. Other specialists are required as practices expand.')
 h(d,'Participation terms',2)
 p(d,'Use defined retainers for availability and review, scoped fees for capture or teaching, and licensing arrangements for content. Specify scope, turnaround, workload, attribution, conflicts, quality obligations and continuity. Royalties must define their revenue base, deductions, timing and audit rights. Equity is a separate parent decision and is not assumed in launch ownership.')
 table(d,['Asset or permission','Responsibility and proposed treatment'],[
  ('Expert background content','The expert or actual rights holder retains pre-existing ownership. The subsidiary secures explicit licenses for permitted uses. Authorship alone does not establish publisher or former-client permissions.'),
  ('New cases and teaching material','The contract specifies assignment or license, attribution, compensation, revisions and post-termination use. Technical Director owns acceptance; Knowledge Engineer keeps asset and rights records.'),
  ('Software and improvements','Head of Product & AI maintains the IP register. New subsidiary-funded software should be assigned to the subsidiary through employment and contractor agreements. Parent or third-party components need documented licenses.'),
  ('Customer records and derived cases','The customer retains its contractual rights in supplied data. Managing Director approves permitted uses on legal advice. Delivery access does not permit general model training, cross-customer reuse or publication.'),
  ('Rights evidence and controls','Knowledge Engineer records source, rights holder, grant, restrictions, expiry and use history. Product leadership enforces access and removal. Unclear material is held out until cleared.')],[2,5.2])
 p(d,'Contracts separately address retrieval, training or fine-tuning, derived content, redistribution and sublicensing. Removing names alone does not establish reuse permission. Material IP assignments and licenses remain parent reserved matters; operational licensing within delegation remains with the Managing Director.')
 page(d,'5 Customer delivery and product acceptance')
 table(d,['Step','Accountable owner','Required result'],[
  ('1 Qualify','Commercial Director','Named buyer and sponsor, recurring problem, procurement route and scope fit.'),
  ('2 Accept technical scope','Technical Director','Supported domain, required experts, assumptions, exclusions and review plan.'),
  ('3 Agree engagement','Managing Director','Signed scope, price, acceptance criteria, data permissions and support limits within delegation.'),
  ('4 Mobilize and baseline','Delivery owner','Team, customer access approval, schedule and baseline. This owner is the Technical Director at launch and the Delivery Lead after appointment.'),
  ('5 Prepare and build','Head of Product & AI','Permitted sources, reviewed cases, reproducible tools, separated evaluation data and controlled implementation.'),
  ('6 Accept technical output','Technical Director','Independent findings resolved; evidence, calculations, limitations and escalation behavior accepted.'),
  ('7 Authorize product release','Head of Product & AI','Technical and security acceptance recorded, deployment conditions met, monitoring and rollback ready.'),
  ('8 Deliver and obtain acceptance','Delivery owner','Customer sign-off or documented exceptions, training, support handover and actual effort recorded.'),
  ('9 Renew or extend','Managing Director','Commercial Director recommends terms using adoption, value, cost and fit. Changed scope returns to technical acceptance.')],[1.6,1.8,3.8])
 h(d,'Evidence required for technical acceptance',2)
 p(d,'Start with retrieval of cleared sources, engineering tools and expert review. Train or fine-tune models only when comparative evaluation demonstrates a need. Keep evaluation incidents separate from development material. Record provenance, units, assumptions, boundary conditions and when the system must abstain or escalate.')
 p(d,'Independent reviewers must be competent in the domain and independent of the material being certified. If the Technical Director authored the work, assign a separate reviewer; the Technical Director remains accountable for resolving findings and accepting the deliverable. No author may self-certify an independent review.')
 p(d,'Customers retain responsibility for plant operating decisions and their authorization procedures. Outputs identify intended use and limits. Any service requiring professional authorization is scoped with parent legal support before sale. Critical technical findings trigger suspension by the Technical Director, correction, review and reacceptance.')
 page(d,'6 Shared services and operating economics')
 p(d,'Broadbridge4096 supplies shared services through a written arrangement with named owners, service levels, cost allocation, access boundaries and escalation routes. The Managing Director is accountable for the service agreement and operating budget. Customer delivery and product priorities remain with the subsidiary.')
 table(d,['Parent service','Service responsibility','Subsidiary interface'],[
  ('Finance','Accounting, reporting, billing support and financial controls.','Managing Director owns P&L; Commercial Director supplies accepted commercial terms.'),
  ('Treasury','Cash planning, approved funding and banking support.','Managing Director forecasts cash and requests funding through the parent capital process.'),
  ('Legal','Contracts, IP, professional-service requirements and legal implementation.','Managing Director owns contract decisions; Technical Director defines service scope.'),
  ('HR','Recruitment, employment administration, compensation and policies.','Managing Director owns workforce planning; parent approves reserved senior appointments.'),
  ('IT security','Identity, security review, monitoring standards and incident response support.','Product leadership implements controls; designated security lead accepts them.'),
  ('Corporate administration','Board records, statutory administration and entity records after formation.','Managing Director supplies decisions and records; board owns oversight.')],[1.2,2.7,3.3])
 h(d,'Revenue and cost responsibility',2)
 p(d,'Maintain separate revenue and delivery cost records for advisory, subscriptions, training and licensing. Track expert hours, royalties, cloud usage, implementation and parent service charges. Allocate shared costs on an agreed basis so practice margins and subsidiary results remain interpretable.')
 p(d,'The Commercial Director recommends pricing; the Managing Director approves commitments within delegation. Product leadership estimates implementation and support; the Technical Director estimates expert and review capacity. Renewal pricing must reflect actual service consumption.')
 h(d,'Financial assumptions',2)
 p(d,'The staffing stages are capacity designs, not cost estimates or approved hiring budgets. Salaries, benefits, expert retainers, royalties, parent charges, software costs and working capital remain to be priced. No revenue, margin, salary or funding forecast is adopted by this document.')
 p(d,'The companion assessment contains illustrative pricing and unit economics. Its bounded validation expenditure hypothesis is not a fully costed budget for seven employees. Parent funding approval requires a cash plan with hiring dates, compensation assumptions and downside runway.')
 page(d,'7 Staffing stages and practice expansion')
 staffing(d)
 p(d,'Employees are counted once. Contractors and parent shared-service personnel are excluded. Technical Director is included in technical practices; Head of Product & AI is included in product, AI and knowledge engineering.')
 table(d,['Function','Expansion role mix','Mature role mix'],[
  ('Managing Director','1 Managing Director','1 Managing Director'),
  ('Commercial','Commercial Director + 1 Account Manager','Commercial Director + 2 Account Managers'),
  ('Technical practices','Technical Director + Downstream Practice Lead + Gas Practice Lead + Process Applications Engineer','Technical Director + 3 Practice Leads + 3 Process Applications Engineers'),
  ('Product, AI and knowledge','Head of Product & AI + Applied AI Engineer + Knowledge Engineer + Platform Engineer','Head of Product & AI + 2 Applied AI Engineers + 2 Knowledge Engineers + Platform Engineer'),
  ('Delivery and customer success','Delivery Lead + Customer Success Manager','Delivery Lead + 2 Delivery Engineers + Customer Success Manager'),
  ('Academy','No dedicated employees','Head of Academy + Training Designer'),
  ('Business operations','No dedicated employee','Business Operations Manager')],[1.5,2.8,2.9],10)
 p(d,'The Managing Director manages the three launch leaders and, when appointed, the Delivery Lead, Head of Academy and Business Operations Manager. Practice leads report to the Technical Director. At expansion the Process Applications Engineer reports to the Downstream Practice Lead; at maturity one reports to each practice lead. Delivery Engineers coordinate implementation and are distinct from practice engineering positions.')
 p(d,'Account Managers report to the Commercial Director; product, AI and knowledge positions to the Head of Product & AI; delivery and customer success to the Delivery Lead; Training Designer to Head of Academy. Technical Director retains technical acceptance across all teams.')
 h(d,'Conditions for the next staffing stage',2)
 p(d,'The Managing Director presents evidence of repeatable paid deployments and funded staffing to the board; the parent approves the budget. Each new practice needs a paid mandate, specialists, cleared knowledge and demonstrated technical capability. Appointments proceed against approved positions. The 24-person structure is a mature target, not a launch commitment.')
 page(d,'8 Performance and implementation')
 table(d,['Measure','Accountable owner','Review cadence'],[
  ('Qualified paid pipeline, win rate and renewal conversion','Commercial Director','Monthly; separate signed commitments from interest.'),
  ('Revenue and contribution by offering, expert hours per site and cash runway','Managing Director','Monthly; finance validates calculations and parent charges.'),
  ('Review findings, unsupported claims, critical failures and escalation quality','Technical Director','Each acceptance and monthly trends.'),
  ('On-time delivery, acceptance and agreed workflow value','Delivery owner','Each engagement and monthly; baseline and sign-off substantiate value.'),
  ('Evaluation quality, use, availability and implementation effort','Head of Product & AI','Each release and monthly; consistent comparison baseline.'),
  ('Rights coverage, expired grants and provenance','Knowledge Engineer','Before ingestion/release and monthly; unresolved rights block use.'),
  ('Training completion, outcomes and instructor capacity','Technical Director at launch; Head of Academy after appointment','Each program; ownership transfers on appointment.'),
  ('Security exceptions and incident response','Designated parent IT security lead','At release, after incidents and through service review.')],[3,1.65,2.55],10)
 p(d,'Set commercial and operating targets after baseline discovery and pilot scope agreement. From first use, every deployed source needs recorded permissions and every release needs required acceptance. Unresolved critical technical or security findings block release. Passing a finite evaluation does not establish universal safety or coverage.')
 h(d,'First ninety days',2)
 table(d,['Period','Implementation priority'],[
  ('Days 1 to 15','Managing Director secures delegated authority and a costed staffing plan. Technical Director confirms expert scope and rights route. Commercial Director qualifies paid pilot demand.'),
  ('Days 16 to 40','Knowledge Engineer and experts capture permitted cases. Product leadership builds a bounded workflow. Delivery owner agrees baseline, acceptance criteria and reviewers.'),
  ('Days 41 to 65','Technical Director commissions independent evaluation on reserved cases. Product leadership fixes failures and completes security acceptance before deployment.'),
  ('Days 66 to 90','Delivery owner runs the limited pilot. Commercial Director develops the renewal case. Managing Director presents results, costs and the next funding recommendation.')],[1.15,6.05])
 p(d,'Timing depends on experts, procurement and data permissions. Legal implementation resolves jurisdiction, final legal form and required service authorizations. Hiring, expert contracts and capital commitments follow the applicable approvals. Candidate names and planning assumptions are not executed appointments or funding.')
 d.save(OUT/'Broadbridge_Oil_and_Gas_Business_and_Operating_Structure.docx')

if __name__=='__main__':
 chart(); operating()
 print('Created 2 editable Word documents. Staffing totals: 7 / 13 / 24.')
