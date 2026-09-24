"""Build the editable handoff document from its Markdown source."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output/foundry-infrastructure'
SOURCE = OUT / 'SLM_Foundry_Broadbridge_Infrastructure_Brief.md'
DEST = SOURCE.with_suffix('.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.6)
sec.left_margin = sec.right_margin = Inches(.65)
sec.footer_distance = Inches(.25)
styles = doc.styles
for name in ['Normal', 'Body Text', 'List Bullet', 'List Number']:
    s = styles[name]
    s.font.name = 'Calibri'
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.0
styles['Title'].font.name = 'Calibri'
styles['Title'].font.size = Pt(25)
styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
styles['Title'].paragraph_format.space_after = Pt(8)
styles['Heading 1'].font.name = 'Calibri'
styles['Heading 1'].font.size = Pt(18)
styles['Heading 1'].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 1'].paragraph_format.space_before = Pt(0)
styles['Heading 1'].paragraph_format.space_after = Pt(9)
styles['Heading 1'].paragraph_format.keep_with_next = True

footer = sec.footer.paragraphs[0]
footer.alignment = 2
r = footer.add_run('Broadbridge Oil & Gas  |  ')
r.font.size = Pt(8)
r.font.color.rgb = RGBColor.from_string('666666')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def inline(p, text):
    last = 0
    for match in re.finditer(r'\[([^\]]+)\]\((https?://[^)]+)\)', text):
        p.add_run(text[last:match.start()])
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(match.group(2), RT.HYPERLINK, is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        color = OxmlElement('w:color'); color.set(qn('w:val'), '146082'); props.append(color)
        run.append(props)
        t = OxmlElement('w:t'); t.text = match.group(1); run.append(t)
        link.append(run); p._p.append(link)
        last = match.end()
    p.add_run(text[last:])

def table(lines):
    rows = [[cell.strip() for cell in line.strip().strip('|').split('|')] for line in lines]
    rows.pop(1)
    n = len(rows[0])
    widths = {2: [2.05, 5.15], 3: [1.75, 3.1, 2.35], 7: [2.28, .82, .82, .82, .82, .82, .82]}[n]
    t = doc.add_table(rows=0, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Inches(width)
    for i, vals in enumerate(rows):
        cells = t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        nosplit = OxmlElement('w:cantSplit'); trpr.append(nosplit)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for j, (cell, value) in enumerate(zip(cells, vals)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, size in [('top', 70), ('bottom', 70), ('left', 90), ('right', 90)]:
                el = OxmlElement('w:' + side); el.set(qn('w:w'), str(size)); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcpr.append(margins)
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E6EDF1' if i == 0 else ('F4F6F7' if i % 2 == 0 else 'FFFFFF')); tcpr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if n == 7 and j > 0: p.alignment = 1
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(10)
                if i == 0: run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

text = SOURCE.read_text(encoding='utf-8')
lines = text.splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line == '<!-- page -->': doc.add_page_break(); i += 1; continue
    if line.startswith('```'):
        i += 1
        code = []
        while i < len(lines) and not lines[i].startswith('```'):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run('\n'.join(code)); r.font.name = 'Consolas'; r.font.size = Pt(9)
        i += 1; continue
    if line.startswith('|'):
        data = []
        while i < len(lines) and lines[i].startswith('|'): data.append(lines[i]); i += 1
        table(data); continue
    if line.startswith('# '): doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('## '): doc.add_paragraph(line[3:], style='Heading 1')
    elif line.startswith('- '): inline(doc.add_paragraph(style='List Bullet'), line[2:])
    else: inline(doc.add_paragraph(), line)
    i += 1

doc.core_properties.title = 'SLM Foundry infrastructure implementation brief'
doc.core_properties.subject = 'Broadbridge Oil and Gas training infrastructure handoff'
doc.core_properties.author = 'Broadbridge4096'
for border in list(doc.styles.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
for border in list(doc.element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
doc.save(DEST)
print(DEST)
print('Source words:', len(text.split()))
for num, section in enumerate(text.split('<!-- page -->'), 1):
    print('Section', num, 'words', len(section.split()))
