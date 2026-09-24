from pathlib import Path
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt
from build_oil_gas_documents import base, table as word_table

ROOT=Path(__file__).resolve().parents[1]
SOURCES={
 1:('Qwen3 8B official model card','https://huggingface.co/Qwen/Qwen3-8B'),
 2:('Qwen3.5 4B official model card','https://huggingface.co/Qwen/Qwen3.5-4B'),
 3:('Ministral 3 8B Instruct 2512 official model card','https://huggingface.co/mistralai/Ministral-3-8B-Instruct-2512'),
 4:('QLoRA research paper','https://arxiv.org/abs/2305.14314'),
 5:('Hugging Face PEFT quantization guide','https://huggingface.co/docs/peft/developer_guides/quantization'),
 6:('Hugging Face TRL supervised fine tuning','https://huggingface.co/docs/trl/sft_trainer'),
 7:('Hugging Face TRL preference optimization','https://huggingface.co/docs/trl/dpo_trainer'),
 8:('Retrieval augmented generation research paper','https://arxiv.org/abs/2005.11401'),
 9:('vLLM structured outputs documentation','https://docs.vllm.ai/en/latest/features/structured_outputs/'),
 10:('Process Improvement Engineering personnel','https://www.lieberman-eng.com/personnel.htm'),
 11:('Lambda on demand GPU instances','https://lambda.ai/instances'),
}
d=base('Domain specific small language model development plan')
d.paragraphs[2].text='Broadbridge4096 subsidiary program  |  23 September 2026'
d.paragraphs[2].runs[0].font.size=Pt(9)
md=['# Broadbridge Oil & Gas domain specific SLM development plan','23 September 2026','']

def link(p,label,url):
 rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
 e=OxmlElement('w:hyperlink'); e.set(qn('r:id'),rel)
 r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'18364D'); pr.append(color); r.append(pr)
 t=OxmlElement('w:t'); t.text=label; r.append(t); e.append(r); p._p.append(e)

def p(text):
 text=re.sub(r'\$(\d{4,})',lambda m:'$'+format(int(m.group(1)),','),text)
 para=d.add_paragraph()
 for part in re.split(r'(\[\d+\])',text):
  if re.fullmatch(r'\[\d+\]',part): link(para,part,SOURCES[int(part[1:-1])][1])
  else: para.add_run(part)
 for k,(label,url) in SOURCES.items(): text=text.replace(f'[{k}]',f'[{k}]({url})')
 md.extend([text,''])

def h(text,level=1,new=False):
 if new: d.add_page_break()
 d.add_heading(text,level); md.extend([('#'*(level+1))+' '+text,''])

def table(headers,rows,widths,size=10.5):
 rows=[[re.sub(r'\$(\d{4,})',lambda m:'$'+format(int(m.group(1)),','),v) if isinstance(v,str) else v for v in row] for row in rows]
 word_table(d,headers,rows,widths,size)
 md.append('| '+' | '.join(headers)+' |'); md.append('|'+'|'.join(['---']*len(headers))+'|')
 for row in rows: md.append('| '+' | '.join(str(c).replace('\n','; ') for c in row)+' |')
 md.append('')

h('1 Recommended program')
p('Build a Broadbridge-owned domain adapter on a commercially usable open-weight small language model, supported by licensed retrieval and deterministic engineering tools. Launch within Downstream and Petrochemicals, beginning with distillation and vacuum-system troubleshooting and case-based training. Use the working product name Broadbridge Process SLM. Broadbridge Oil & Gas remains a subsidiary of Broadbridge4096.')
p('Fine-tuning aims to teach the model to assemble evidence, ask useful diagnostic questions, use tools, cite retrieved material and escalate uncertainty. Current plant facts, customer records and licensed reference passages belong in permission-controlled retrieval. A trained model is one component of the customer product, not a substitute for the engineering review workflow.')
table(['Planning decision','Recommended starting point'],[
 ('Model size','Compare a roughly 4B candidate with an 8B-class candidate; select on Broadbridge evaluation and measured operating cost.'),
 ('Training method','Supervised fine-tuning using LoRA or QLoRA adapters. Optional preference training follows only if it fixes measured weaknesses.'),
 ('Deployment','Private cloud first, with a later on-premises option. This is a provisional assumption pending customer requirements.'),
 ('Schedule','24 weeks to a restricted paid pilot, conditional on data rights, staffing and customer access. A first internal adapter is targeted by week 12.'),
 ('Initial funding decision','Approve only the first eight weeks and its deliverables, then decide whether the evidence supports the next stage.')],[1.7,5.5])
h('Approaches considered',2)
table(['Approach','Tradeoff and decision'],[
 ('Retrieval with an untuned model','Fast baseline and a viable product if tuning adds little. Keep it as a control and fallback.'),
 ('Fine-tuned SLM plus retrieval and tools','Recommended. Adds a repeatable domain workflow while keeping facts and permissions outside the weights.'),
 ('Domain pretraining or training from scratch','Defer. Requires a different data and compute case; first determine whether supervised adaptation addresses the observed gaps.')],[2.0,5.2])
p('The existing assessment remains the commercial starting point. The Energy Trading materials are organizational references only. Norman Lieberman is an illustrative expert candidate, not an appointed contributor. This plan authorizes neither data use nor model deployment; those steps follow the subsidiary decision rights.')

h('2 Product behavior and architecture',new=True)
p('The first user is a process engineer preparing a diagnostic brief for expert review. Start with text, structured operating snapshots and approved historical cases. Validate digitized tables and units before use. Automated interpretation of process drawings, live plant control, drilling and subsurface reasoning are outside the first release.')
h('Runtime sequence',2)
p('Authenticated user → scope and tenant check → retrieval of authorized sources → SLM with selected context → approved calculation tools → claim and schema checks → diagnostic brief → engineer or expert review.')
table(['Component','Responsibility and interface'],[
 ('Knowledge service','Search versioned passages with source IDs, page references, dates and permissions. Filter by customer and user authorization before retrieval and reranking.'),
 ('Broadbridge domain adapter','Apply the reviewed diagnostic method. Return hypotheses, evidence gaps, tool requests, citations and escalation status in a defined output schema.'),
 ('Engineering tools','Accept typed inputs and units, validate ranges and applicability, and return calculations plus assumptions. The application allows only named tools; no unrestricted code execution.'),
 ('Application checks','Verify that cited IDs exist in the retrieved set, units are present and required fields are valid. Semantic support still requires evaluation and technical review.'),
 ('Review and audit','Capture model/adapter version, permitted sources, tool results, review decisions and corrections under a defined retention policy.')],[1.65,5.55])
h('Diagnostic brief output',2)
p('Require the problem statement; equipment and operating context; known facts and their sources; competing hypotheses; supporting and conflicting evidence; missing measurements; checked calculations; assumptions; limitations; and escalation or review requirements. The response is a reviewable assessment, not an instruction to change plant limits or equipment settings.')
p('When evidence is missing, the model asks a focused question or states that it cannot support a diagnosis. It must not invent a source, measured value or confident numerical probability. Training should use concise expert explanations with visible evidence and calculations; hidden model reasoning is not an audit record.')
p('Retrieval gives access to external evidence; fine-tuning changes model behavior. The research basis supports combining them, but does not establish performance for refinery operations. Structured output support can enforce form, not engineering correctness. [8] [9]')

h('3 Foundation model selection',new=True)
p('Begin with Qwen3-8B as the reference training candidate and Qwen3.5-4B as the efficiency challenger. Include Ministral 3 8B as an alternative if customer policy or measured results favor it. These are test candidates, not a vendor award. Public benchmarks do not measure the proposed Broadbridge workflow.')
table(['Candidate','Verified model facts','Selection issue'],[
 ('Qwen3-8B','8.2B parameters; Apache 2.0 model card.','Reference for the text workflow. Test tool calls, citation grounding, domain errors and inference cost.'),
 ('Qwen3.5-4B','4B language model with vision encoder; Apache 2.0 model card.','Efficiency challenger. Verify hybrid-architecture training and serving support; vision is not in launch scope.'),
 ('Ministral 3 8B Instruct 2512','8.4B language model plus 0.4B vision encoder; Apache 2.0. Published instruct weights use FP8.','Check training-compatible precision and adapter support. Do not assume FP8 inference weights can directly use the reference QLoRA recipe.')],[1.65,2.65,2.9],10.5)
p('The sizes, license labels and format details above come from the publishers. The recommendation to test them is Broadbridge program design, not a claim that one is best in oil and gas. Confirm the exact revision, full license and required notices before adoption. [1] [2] [3]')
h('Selection experiment',2)
p('Run all candidates on the same development questions, retrieved evidence, tool definitions and output requirements. Allow model-specific chat templates and a documented equal tuning allowance. Measure task acceptance, missing-data handling, calculation/tool success, latency, memory and total cost. Use a current larger model as a quality comparator only in an environment authorized for the data.')
p('Select on the development set before opening the locked acceptance set. Prefer the smallest model that meets the technical requirements and materially improves total delivery economics. A larger model may remain a comparator; external fallback must be explicitly enabled for a customer, never a hidden data route.')
h('Compatibility gate before training',2)
p('Pin the model and tokenizer revision, training framework, quantization library, GPU software and inference engine. Prove a small adapter can train, save, reload and serve without changing chat behavior. Verify the output schema and tool parser end to end. If the challenger needs unsupported kernels or experimental conversion, retain the reference candidate and record the constraint.')
p('Use English for the first release. Additional languages, longer context, new hardware and additional practices each require evaluation. Published maximum context windows are not the initial application setting; begin with a bounded evidence budget and measure what the task needs.')

h('4 Expert capture and data rights',new=True)
p('The durable asset is a permitted, reviewed set of diagnostic cases with outcomes and limitations. Published books and articles alone are insufficient: extract the questions an experienced engineer asks, the observations that distinguish causes, the calculations that matter, and the conditions that require a different specialist.')
p('Recruit a lead process expert and a separate reviewer, supported by the Knowledge Engineer and Process Applications Engineer. Norman Lieberman is an illustrative candidate based on his published practice and teaching background; availability, endorsement and rights are unconfirmed. Recruit complementary contributors rather than making the program depend on one individual. [10]')
h('Capture workflow',2)
p('For each incident, record the initial evidence, competing explanations, discriminating checks, assumptions, outcome and circumstances where the approach would fail. Ask the expert to contrast a similar-looking incident with a different cause. Preserve disagreements for adjudication rather than forcing a single unsupported answer. Obtain permission before recording or transcribing interviews.')
table(['Material','Permitted role in the program'],[
 ('Expert interviews and commissioned cases','Use after an executed agreement specifies recording, derivative cases, training, evaluation, distribution, attribution and continuity rights.'),
 ('Books, articles and standards','License the intended uses from the actual rights holders. Public access or author participation does not establish publisher permission.'),
 ('Customer incident histories and records','Use for that customer only under its agreement. Shared model training or cross-customer examples require a separate explicit grant.'),
 ('Public technical material','Verify license and provenance; review applicability and quality. Use as a supplement, not an automatic source of training rights.'),
 ('Synthetic examples','Generate only from cleared sources or expert-defined scenarios. Label provenance and expert approval; never treat a paraphrase as a new independent incident.')],[1.85,5.35])
h('Separate ownership and permissions',2)
p('Broadbridge Oil & Gas should own commissioned adapters, training code and new datasets to the extent established by its agreements. The foundation model remains subject to its original license. Expert background works retain their ownership unless assigned. A customer-data grant must separately address retrieval, training, evaluation, publication and retention.')
p('Maintain a rights ledger with source ID, rights holder, contract, allowed uses, customer boundary, expiry, withdrawal terms and derivative lineage. Train shared adapters only on material cleared for that purpose. Prefer retrieval for revocable or customer-specific information: deleting a source from a library does not reliably remove information already learned by an adapter.')

h('5 Dataset and benchmark construction',new=True)
p('Planning target by week 16: 300 distinct incident or scenario families, split before examples are generated. This is a collection target, not an inventory already available or a universal minimum for successful fine-tuning. If rights or case diversity fall short, narrow the claim and extend collection rather than manufacture volume.')
table(['Partition','Target','Use and restriction'],[
 ('Training','180 families','Create the supervised training examples. Related incidents, source chapters, derivatives and paraphrases stay together.'),
 ('Development','60 families','Select models, prompts, retrieval settings and training parameters. Never describe this tuned-against set as an independent final test.'),
 ('Locked acceptance','60 families','Reviewer-controlled until candidate selection is frozen. Two tasks per family produce 120 diagnostic prompts.'),
 ('Additional challenge suite','180 prompts','Separately authored tests of missing data, unit errors, unsupported scope, tool misuse, prompt injection and data boundaries. Track scenario families and correlation.')],[1.5,1.15,4.55])
p('The acceptance suite contains 300 prompts, not 300 independent field incidents. Group results by incident/scenario family. Exclude held-out resolutions and answer keys from training, retrieval and tool outputs; supply only evidence available at question time. Add an unseen-site or later-time slice where feasible. Report public teaching cases separately because their absence from foundation-model pretraining cannot be guaranteed.')
h('First full supervised dataset',2)
table(['Behavior','Examples in a 2000 record target'],[
 ('Evidence-based diagnostic briefs',900),
 ('Missing-data questions and differential diagnosis',400),
 ('Tool use, units and calculation interpretation',300),
 ('Source-grounded explanation and training feedback',200),
 ('Abstention, escalation and unsupported scope',200)],[5.2,2.0])
p('Begin learning-curve experiments with 500 approved records, then 1000 and 2000; expand toward 4000 only when additional diversity improves validation. All records receive technical review. Generation can accelerate drafting but cannot replace review. Multiple records from one family increase practice examples, not independent evidence.')
p('Each record includes record ID, parent family ID, source and rights IDs, practice, equipment, input evidence, units, intended task, permitted context, tool inputs/results, approved response, limitations, author, reviewer, dates and split. Block unlicensed records, duplicate families across splits, missing units and unresolved reviewer findings. Separate ingestion mistakes from model errors.')

h('6 Fine tuning and experiment program',new=True)
p('Use supervised fine-tuning first: train the model against expert-approved responses and tool-use sequences. LoRA updates a relatively small set of adapter parameters; QLoRA combines adapter training with a quantized foundation model to reduce memory needs. Neither technique establishes factual accuracy or confidentiality by itself. [4]')
h('Reference starting configuration',2)
table(['Setting','Proposed experiment value'],[
 ('Reference foundation','Pinned Qwen3-8B instruction/post-trained checkpoint; preserve its native chat template.'),
 ('Adapter and precision','QLoRA with 4-bit NF4 frozen weights and BF16 computation where supported; compare with BF16 LoRA on a small run.'),
 ('Adapter capacity','Start rank 16, alpha 32, dropout 0.05; compare rank 32 only if validation supports more capacity.'),
 ('Optimization','Start learning rate 1e-4, effective batch 32 and one epoch; test at most three epochs with validation-based stopping.'),
 ('Sequence length','Start at 4096 tokens; test 8192 where evidence is truncated. Profile memory and throughput before enlarging.'),
 ('Targets and loss','Select supported language projection layers after architecture inspection. Train on approved assistant/tool-call targets and mask prompt/context loss; verify masking and packing boundaries.'),
 ('Reproducibility','Keep dataset hashes, split manifests, seed, configuration, dependencies and run metrics. Repeat the selected configuration across three seeds.')],[1.7,5.5],10.5)
p('These values are starting hypotheses, not an optimized recipe or a promise of hardware fit. Hugging Face documents the quantized-adapter and supervised-training mechanisms; architecture-specific compatibility must be tested. [5] [6]')
h('Experiments that isolate the contribution',2)
p('Compare A: untuned SLM with retrieval and tools; B: tuned SLM without retrieval as a diagnostic ablation; C: tuned SLM with the same retrieval and tools; and D: a larger approved model with the same evidence and tools. A versus C measures tuning benefit. B is not a deployment candidate. Compare on unchanged evaluation inputs and separately report retrieval failures.')
p('Use preference optimization only if supervised training leaves consistent, reviewable weaknesses. Collect at least 300 adjudicated preferred/rejected response pairs as an experiment target, then compare against supervised training alone. Do not use synthetic self-preference as proof of quality. [7]')
p('Defer continued domain pretraining until learning curves reveal a gap that retrieval and supervised training cannot address and a licensed corpus and separate budget exist. Defer training from scratch. Treat deployment quantization as a separate export decision from QLoRA training. Repeat acceptance checks on the exact exported model, adapter and serving configuration.')

h('7 Evaluation and release criteria',new=True)
p('The Technical Director owns technical acceptance. A reviewer who did not author the evaluated material controls the locked test and assesses blinded outputs. Use at least two independent ratings for material diagnostic cases, with disagreement adjudication. Test the complete product and the selected quantized artifact, not just the training checkpoint.')
table(['Measure','Proposed acceptance rule'],[
 ('Useful diagnostic output','At least 85% of applicable held-out diagnostic tasks accepted without a material technical correction, using a rubric agreed before test access.'),
 ('Benefit from fine-tuning','Target at least a 10 percentage-point improvement over A, the untuned model with identical retrieval/tools. Report paired confidence intervals clustered by family; inconclusive results require more evidence.'),
 ('Grounding','At least 95% of sampled material factual claims supported by supplied sources or verified tool results; zero fabricated citation IDs. Record denominator and reviewer sampling method.'),
 ('Calculations and tools','At least 98% correct on the predefined numeric/tool tasks within expert-set tolerances. Any critical error still blocks release.'),
 ('Uncertainty and scope','At least 95% appropriate abstention or escalation on unanswerable/out-of-scope tasks. Report false refusals separately so indiscriminate refusal cannot pass.'),
 ('Critical failures','Zero unresolved critical unsafe recommendations, cross-customer disclosures or successful unauthorized tool actions in the acceptance suite.'),
 ('Runtime','Provisional p95 complete-brief latency of 30 seconds at five concurrent users, for an 8000-token input and up to 800 output tokens. Confirm feasibility with the customer.'),
 ('Customer value','Target 25% less time to an expert-accepted brief than the agreed current workflow, without worse quality. Measure expert minutes and cost per accepted brief.')],[1.65,5.55],10.5)
p('These are proposed program gates, not achieved scores or industry standards. A critical failure overrides an average pass rate. A finite suite cannot prove operational safety or absence of data leakage. Report scenario coverage, uncertainty and the effect of correlated examples; never treat 300 prompts as 300 independent trials.')
p('Use the locked set once for the release decision after development is frozen. If failures drive development changes, retire exposed cases from independence claims and commission a fresh holdout. Technical Director accepts domain behavior; parent IT security accepts security; Head of Product & AI authorizes release only after both acceptances.')
p('If tuning does not add a measurable benefit, retain the retrieval-based product and continue expert capture. If a smaller model meets quality but cannot meet latency, adjust hardware or scope and retest. Do not weaken technical thresholds merely to label the outcome a fine-tuned product.')

h('8 Delivery schedule and gates',new=True)
p('Allow 24 weeks from funded staffing and executed data access. Expert capture, engineering and customer procurement overlap. Delayed rights or insufficient unseen cases move the schedule. The following are deliverables with evidence requirements, not automatic calendar approvals.')
table(['Period','Accountable lead','Deliverable and gate'],[
 ('Weeks 1 to 4','Managing Director','Paid pilot problem and sponsor; contributor terms; rights inventory; scope and output schema; named reviewers. Freeze split policy and reserve unseen sources before generating examples. Parent approves stage funding.'),
 ('Weeks 5 to 8','Head of Product & AI','Candidate benchmark, retrieval/tool baseline, roughly 50–60 cleared training families plus a separate development sample; first 500 reviewed training records where available. Parent decides the next funding tranche.'),
 ('Weeks 9 to 12','Head of Product & AI','First supervised domain adapter; 500/1000-record learning curves; reproducible save/reload; comparison with untuned baseline. Continue expert collection.'),
 ('Weeks 13 to 16','Technical Director','Target corpus of 300 families and 2000 reviewed training records; independent acceptance set sealed. Select training configuration on development data and document limitations.'),
 ('Weeks 17 to 20','Head of Product & AI','Freeze candidate and serving package; execute independent technical tests, security tests, load tests and rollback rehearsal. Technical and security owners provide separate acceptance.'),
 ('Weeks 21 to 24','Technical Director','Limited paid pilot with 5–10 engineers at 1–2 sites; baseline comparison, human review, incident monitoring and cost recording. Managing Director decides the commercial rollout recommendation.')],[1.1,1.6,4.5],10.5)
h('Required artifacts at the pilot gate',2)
p('Deliver the foundation revision and license record; trained adapter and checksums; inference package; permitted corpus snapshot and rights ledger; training/development/test manifests; evaluation report with limitations; model card; source-citation and tool schemas; security acceptance; user instructions; rollback runbook; and the signed pilot scope. Keep restricted customer data in its approved storage rather than bundling it into a general release.')
h('Engineering work packages',2)
p('Organize implementation into independently reviewable packages: data ingestion and rights enforcement; retrieval and source provenance; diagnostic output and engineering tools; supervised training; evaluation; and deployment/monitoring. Each package must expose versioned inputs and outputs. Keep the locked benchmark access separate from the training pipeline.')
p('Before commissioning the full code build, the product lead translates this program into repository tasks using the selected model and deployment environment. Acceptance must include revoked-source exclusion, split isolation, incorrect-unit rejection, adapter reload equivalence and recovery to the previous approved release.')

h('9 Team allocation and planning budget',new=True)
p('Use the approved seven-person launch organization. The allocations below are program capacity assumptions across 24 weeks, not additional hires. Parent shared services and contracted experts are separate. Technical staff need protected time; simultaneous customer work reduces available capacity and extends the schedule.')
table(['Launch role','Program FTE','Hours over 24 weeks'],[
 ('Managing Director',0.25,240),('Commercial Director',0.25,240),('Technical Director',0.50,480),
 ('Process Applications Engineer',0.75,720),('Head of Product & AI',0.50,480),
 ('Applied AI Engineer',1.00,960),('Knowledge Engineer',1.00,960),('Total allocated capacity',4.25,4080)],[4.35,1.25,1.6],10.5)
p('The Applied AI Engineer leads training and deployment; the Knowledge Engineer leads corpus production; the Process Applications Engineer checks engineering content and tools. The Technical Director manages experts and acceptance. The Commercial Director recruits the pilot; the Managing Director owns funding and contracts. Reserve 600–900 contracted expert/reviewer hours across capture, annotation, review and evaluation.')
table(['Cost category','Illustrative basis','Range USD'],[
 ('Allocated employee cost','4080 hours × $100–$150 loaded/hour','$408000–$612000'),
 ('Experts and independent review','600–900 hours × $250–$350/hour','$150000–$315000'),
 ('Compute and supporting cloud','Training, evaluation, pilot serving, storage and monitoring allowance','$10000–$25000'),
 ('Legal and security support','Incremental external work or parent service allocation; count once','$25000–$50000'),
 ('Content licensing','Provisional allowance for negotiated rights, not a publisher quote','$20000–$60000'),
 ('Subtotal','Before contingency','$613000–$1062000'),
 ('Contingency','15% of subtotal','$91950–$159300'),
 ('Total program economic cost','Allocated labor plus other costs','$704950–$1221300')],[2.0,3.55,1.65],10)
p('These are planning assumptions, not market quotes, an approved budget or a valuation. If employee payroll is already funded, the nonemployee categories plus 15% contingency are approximately $236000–$518000, subject to which parent services are already covered. Do not add allocated payroll twice. The balance of seven-person payroll and other subsidiary costs sits outside this program allocation.')
p('First eight-week tranche: approximately $210000–$360000 including allocated staff, or $54000–$124000 beyond funded staff under the stated assumptions. It is part of the 24-week total, not an additional budget. Obtain quotes and a monthly cash plan before approval; scarce content rights can exceed the allowance.')

h('10 Infrastructure and operating controls',new=True)
p('Start with rented GPU capacity in a customer-approved private environment. Plan experiments around one 48–80 GB GPU and test a 24–48 GB inference GPU for the restricted workload. These are sizing allowances, not verified fit. Memory depends on architecture, sequence length, adapter, precision, batch size and concurrent requests.')
p('For orientation, 200–800 training/evaluation GPU-hours at an assumed $4–$8 per hour cost $800–$6400. Pilot serving, storage, networking and monitoring are additional. Small-model training can be affordable while the complete expert-reviewed program remains labor-intensive. Obtain current regional GPU quotes and verify availability; the provider page is a procurement reference, not the source of these assumed rates. [11]')
h('Deployment controls',2)
p('Use a private model registry, encrypted storage, named identities and least-privilege service accounts. Keep customer retrieval indexes and permissions isolated. Turn off external telemetry or content logging unless approved, and test outbound-network restrictions. A private network label alone does not prove confidentiality.')
p('Version the model, adapter, corpus, prompt, tools and schema together. Monitor unsupported claims, expert corrections, retrieval misses, stale sources, refusal behavior, latency and cost. Corrections enter a reviewed dataset queue; customer feedback must not trigger automatic online training. Retest before each release and retain the prior approved package for rollback.')
p('Treat retrieved documents as evidence rather than instructions. Test prompt injection, invented sources, expired licenses, cross-customer requests and malicious tool arguments. Access enforcement lives in application code and storage policies, not the model prompt. Human review remains required for technical advice in the initial pilot.')
h('Main risks and responses',2)
table(['Risk','Program response'],[
 ('Expert bottleneck or unavailable candidate','Contract multiple contributors, preserve capture methods and schedule review time before committing the pilot date.'),
 ('Small or repetitive training corpus','Measure unique families and learning curves; narrow scope when diversity is insufficient.'),
 ('No improvement after fine-tuning','Keep the untuned retrieval baseline deployable and redirect effort to data, retrieval or the workflow.'),
 ('Rights withdrawal or customer leakage','Keep restricted information out of shared weights; trace affected artifacts and rebuild an adapter if required.'),
 ('On-premises requirement emerges','Reprice hardware, support and installation; benchmark the exact offline package before promising parity.')],[2.0,5.2])
p('The next parent decision is the eight-week discovery and baseline tranche. Before training, confirm the expert/rights package, paid customer problem, cloud versus on-premises constraint, permitted foundation models, evaluation rules and actual staff availability. Expand later to gas processing or upstream only through separately staffed, funded and evaluated domain modules.')

h('11 Research references and program records',new=True)
p('Public technical sources reviewed on 23 September 2026. Model cards and library documentation can change; preserve the exact versions and license files selected during implementation. These sources support the technical methods and model facts, not Broadbridge-specific performance or cost projections.')
for k,(label,url) in SOURCES.items():
 para=d.add_paragraph(); para.add_run(f'{k}  '); link(para,label,url)
 md.extend([f'{k}. [{label}]({url})',''])
h('Broadbridge records',2)
p('The approved subsidiary organization supplies the seven roles, parent reserved decisions, technical acceptance authority and phased practices. The existing business assessment supplies the refining-first commercial hypothesis and the distinction between model training and expert-reviewed product value.')
for label,url in [
 ('Business and operating structure','C:/Users/bradu/Documents/Broadbridge4096/output/docx/Broadbridge_Oil_and_Gas_Business_and_Operating_Structure.docx'),
 ('Subsidiary organizational chart','C:/Users/bradu/Documents/Broadbridge4096/output/docx/Broadbridge_Oil_and_Gas_Organizational_Chart.docx'),
 ('Business assessment','C:/Users/bradu/Documents/Broadbridge4096/Oil-and-Gas-Expert-Knowledge-Business-Assessment.md')]:
 para=d.add_paragraph(); link(para,label,url); md.extend([f'[{label}]({url})',''])
h('Program assumptions requiring confirmation',2)
p('Private-cloud-first deployment; English-only launch; distillation/vacuum focus; availability of cleared incidents; 4.25 FTE allocated from the seven-person team; contracted reviewers; indicative labor and licensing rates; a 24-week funded schedule; and a paid pilot at 1–2 sites. None is a claim of an existing contract, trained model, accepted score or committed customer.')
p('The first eight-week estimate assumes 1360 employee hours at $100–$150/hour; 80–150 expert hours at $250–$350/hour; $15000–$25000 legal/security; $10000–$25000 content rights; and $2000–$5000 cloud. Applying 15% contingency yields $210450–$358225 including labor, or $54050–$123625 excluding funded labor. These rounded tranche ranges are included in the full program estimate.')
p('This plan creates the development roadmap. Broadbridge4096 remains accountable for capital approval; the subsidiary remains accountable for product execution and customer delivery. No content rights, infrastructure purchases or trained-model results are assumed to exist.')

assert 180+60+60==300
assert 900+400+300+200+200==2000
assert 240+240+480+720+480+960+960==4080
assert abs((408000+150000+10000+25000+20000)*1.15-704950)<.01
assert abs((612000+315000+25000+50000+60000)*1.15-1221300)<.01
d.save(ROOT/'output/docx/Broadbridge_Oil_and_Gas_SLM_Development_Plan.docx')
(ROOT/'Broadbridge-Oil-and-Gas-SLM-Development-Plan.md').write_text('\n'.join(md),encoding='utf-8')
print('Created Word and Markdown SLM development plan. Counts and budget reconciled.')
