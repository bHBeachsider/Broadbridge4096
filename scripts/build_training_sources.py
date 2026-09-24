from pathlib import Path
import csv
import json
import re
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Pt
from build_oil_gas_documents import base, table as word_table

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'output'
DATE = '2026-09-23'
records = []

def source(id, name, category, domain, url, rights, rights_url, use, priority, status, owner, action):
    records.append(dict(source_id=id, name=name, category=category, domain=domain,
        source_url=url, observed_rights=rights, rights_evidence_url=rights_url,
        recommended_use=use, priority=priority, disposition=status,
        accountable_owner=owner, next_action=action, reviewed_on=DATE,
        ingestion_status='Not ingested', asset_revision='Capture at acquisition',
        approval_status='No production training approval recorded'))

DOE='https://www.energy.gov/web-policies'
source('P01','DOE Steam Systems and Steam Sourcebook','Public technical reference','Steam and heat transfer',
 'https://www.energy.gov/cmei/ito/steam-systems','Government information generally public domain; contributed material may be protected',DOE,
 'Retrieval; reviewed SFT examples from cleared passages','Launch','Item review', 'Knowledge Engineer','Select sourcebook chapters and tip sheets; inspect authorship and third-party credits.')
source('P02','DOE Process Heating Sourcebook','Public technical reference','Fired heaters and process heat',
 'https://betterbuildingssolutioncenter.energy.gov/better-plants/process-heating','DOE policy with third-party exceptions',DOE,
 'Retrieval; reviewed fundamentals and missing-data examples','Launch','Item review','Knowledge Engineer','Select energy-balance and performance sections; retain edition and limitations.')
source('P03','DOE Pumping System Sourcebook','Public technical reference','Pumps and hydraulic systems',
 'https://betterbuildingssolutioncenter.energy.gov/better-plants/pumps','DOE policy with third-party exceptions',DOE,
 'Retrieval; supporting equipment cases','Launch support','Item review','Knowledge Engineer','Check contributed Hydraulic Institute material separately before reuse.')
source('P04','CSB BP Texas City refinery investigation','Public investigation','Startup evidence and process safety',
 'https://www.csb.gov/bp-america-texas-city-refinery-explosion/','Public report; item and third-party attachment review required','',
 'Historical case reconstruction; escalation examples; public regression tests','Launch','Item review','Technical Director','Use agency findings with dated evidence; separate pre-event information from hindsight and private attachments.')
source('P05','EPA AP 42 Chapter 5 petroleum industry','Public technical reference','Refinery process taxonomy and emissions',
 'https://www.epa.gov/air-emissions-factors-and-quantification/ap-42-fifth-edition-volume-i-chapter-5-petroleum-1',
 'EPA documents can have individual copyright conditions','https://www.epa.gov/web-policies-and-procedures/epa-disclaimers',
 'Versioned retrieval; process vocabulary; emissions context','Launch','Item review','Knowledge Engineer','Review current chapter and credits; keep edition and scope with extracted factors.')
source('P06','EIA Open Data API','Public structured data','Refining capacity throughput and energy context',
 'https://www.eia.gov/opendata/','EIA government products public domain; protected contributed items excepted','https://www.eia.gov/about/copyrights_reuse.php',
 'Live data tool and dated retrieval; not incident diagnosis labels','Context','Item review','Applied AI Engineer','Choose required series only; preserve units frequency revisions and attribution.')
source('P07','CoolProp','Engineering software','Thermophysical properties',
 'https://github.com/CoolProp/CoolProp','MIT software license','https://github.com/CoolProp/CoolProp/blob/master/LICENSE',
 'Checked property tool; generated calculation examples','Launch','Candidate','Applied AI Engineer','Pin version and approved fluids/ranges; validate against independent references; audit optional backends.')
source('P08','IDAES Process Systems Engineering Framework','Engineering software','Balances and process simulation',
 'https://github.com/IDAES/idaes-pse','BSD-style license with additional enhancements provision','https://github.com/IDAES/idaes-pse/blob/main/LICENSE.md',
 'Validated synthetic scenarios and tool traces','Launch','Candidate','Process Applications Engineer','Select a small model; verify solver dependencies model equations and validity before generating labels.')
source('P09','NIST Chemistry WebBook SRD 69','Reference database','Chemical and physical properties',
 'https://webbook.nist.gov/chemistry/','Standard Reference Data may be copyrighted; separate product terms apply',
 'https://www.nist.gov/open/copyright-fair-use-and-licensing-statements-srd-data-software-and-technical-series-publications',
 'Reference validation or licensed lookup; bulk training held pending clearance','Launch support','Permission review','Knowledge Engineer','Resolve specific data reuse and access terms; do not infer public-domain status from .gov.')
source('P10','Equinor Volve field dataset','Public field dataset','Upstream production and subsurface',
 'https://www.equinor.com/energy/volve-data-sharing','Custom Volve terms allow commercial use with attribution; prohibit selling licensed material',
 'https://cdn.equinor.com/files/h61q9gi9/global/de6532f6134b9a953f6c41bac47a0c055a3712d3.pdf?equinor-hrs-terms-and-conditions-for-licence-to-data-volve.pdf=',
 'Later production data tools and separately validated cases','Later upstream','Permission review','Technical Director','Review current package terms and derived commercial outputs before selecting a small subset.')
source('P11','NETL Energy Data eXchange','Data discovery portal','Energy facilities research and later practice data',
 'https://edx.netl.doe.gov/','Collection-level discovery only; individual dataset rights unverified','https://netl.doe.gov/edx',
 'Discover task-specific research data; no blanket ingestion','Later','Discovery','Knowledge Engineer','Select an exact dataset and record its provenance access rights and technical relevance.')

for item in [
 ('M01','Qwen/Qwen3-8B','Foundation model','8B class text model','https://huggingface.co/Qwen/Qwen3-8B','Primary SFT baseline','Launch','Pin revision; compare untuned and QLoRA-tuned versions with the same retrieval and tools.'),
 ('M02','Qwen/Qwen3.5-4B','Foundation model','4B class model with vision capability','https://huggingface.co/Qwen/Qwen3.5-4B','Smaller efficiency challenger; text-only launch','Launch','Validate hybrid architecture adapter targets and serving support; unused vision does not eliminate its implementation overhead.'),
 ('M03','mistralai/Ministral-3-8B-Instruct-2512','Foundation model','8.4B language model plus vision encoder','https://huggingface.co/mistralai/Ministral-3-8B-Instruct-2512','Alternative SFT candidate','Launch alternate','Card describes FP8 instruct checkpoint; verify training format and adapter support before selection.'),
 ('M04','Qwen/Qwen3-32B','Teacher model','General language and instruction','https://huggingface.co/Qwen/Qwen3-32B','Optional local drafting teacher and larger-model comparator','Optional','Use only cleared inputs; independently verify every proposed example and retain teacher metadata.'),
 ('M05','Qwen/Qwen3-Embedding-0.6B','Retrieval model','Passage embedding','https://huggingface.co/Qwen/Qwen3-Embedding-0.6B','Semantic retrieval baseline alongside lexical search','Launch','Measure retrieval recall on engineer-labeled queries; respect tenant filters before retrieval.'),
 ('M06','Qwen/Qwen3-Reranker-0.6B','Retrieval model','Passage ranking','https://huggingface.co/Qwen/Qwen3-Reranker-0.6B','Reorder retrieved evidence','Launch','Measure relevant passage ranking and latency; tune only if labeled evidence justifies it.')]:
    id,name,cat,dom,url,use,pri,action=item
    source(id,name,cat,dom,url,'Apache 2.0 on official model card',url,use,pri,'Candidate','Head of Product & AI',action)

source('M07','PetroGPT/Llama-3-Petro-Instruct-v1','Domain model','Petroleum-labeled instruction model',
 'https://huggingface.co/PetroGPT/Llama-3-Petro-Instruct-v1','Card labels Apache 2.0 but identifies Llama 3 derivative; underlying terms need reconciliation',
 'https://huggingface.co/meta-llama/Meta-Llama-3-8B-Instruct/blob/main/LICENSE',
 'Possible comparator only after model and training-data lineage review','Defer','Hold','Head of Product & AI','Resolve license chain and data provenance; no verified Broadbridge refining score.')
source('M08','GeoGPT-Research-Project/GeoGPT-R1-Preview','Domain model','Geoscience 72B model',
 'https://huggingface.co/GeoGPT-Research-Project/GeoGPT-R1-Preview','Custom GeoGPT license; card specifies noncommercial research and education',
 'https://huggingface.co/GeoGPT-Research-Project/GeoGPT-R1-Preview',
 'Methodology reference; commercial training teacher and deployment held','Later upstream','Hold','Head of Product & AI','Obtain applicable commercial permissions before model use; dataset license is separate.')
source('M09','AI4Chem/ChemLLM-7B-Chat','Domain model','Chemistry and molecular science',
 'https://huggingface.co/AI4Chem/ChemLLM-7B-Chat','Card separates Apache code from weights and requests commercial-license contact',
 'https://huggingface.co/AI4Chem/ChemLLM-7B-Chat','Research comparison; not a refinery baseline','Defer','Hold','Head of Product & AI','If needed later, review current 1.5 variants and the full weights/base-model license chain.')

source('D01','PetroGPT/petro-dataset-v2','Public instruction dataset','Broad chemistry and some domain material',
 'https://huggingface.co/datasets/PetroGPT/petro-dataset-v2','CC BY-NC 4.0 shown on dataset card',
 'https://huggingface.co/datasets/PetroGPT/petro-dataset-v2','Excluded from commercial training and teacher generation without separate rights','Exclude','Hold','Knowledge Engineer','Resolve commercial permission and source provenance before any future use; preview is not narrowly refining-focused.')
source('D02','GeoGPT-Research-Project/GeoGPT-CoT-QA','Public instruction dataset','Geoscience',
 'https://huggingface.co/datasets/GeoGPT-Research-Project/GeoGPT-CoT-QA','CC BY 4.0; card says source publications are CC BY',
 'https://huggingface.co/datasets/GeoGPT-Research-Project/GeoGPT-CoT-QA','Possible later filtered training subset; separate evaluation split','Later upstream','Item review','Knowledge Engineer','Audit DOI licenses and generated answers; use concise checked explanations; limited refining relevance.')
source('D03','petra-2026/PETRA','Retrieval dataset','Petroleum-tagged retrieval and reranking',
 'https://huggingface.co/datasets/petra-2026/PETRA','License marked other; underlying source rights not established in this review',
 'https://huggingface.co/datasets/petra-2026/PETRA','Potential later retrieval experiment; bulk ingestion held','Defer','Hold','Knowledge Engineer','Audit source-level permissions and relevance; displayed sample includes unrelated subject matter.')
source('D04','my2000cup/PetroBench','Public benchmark','Petroleum questions and calculations',
 'https://huggingface.co/datasets/my2000cup/PetroBench','MIT label on dataset card; source lineage requires review',
 'https://huggingface.co/datasets/my2000cup/PetroBench','Secondary benchmark only; excluded from training','Later upstream','Evaluation only','Technical Director','Validate source rights schema encoding equations and answers; freeze a named revision; do not equate with same-named papers without provenance.')
source('D05','FormationEval original 505 question track','Public benchmark','Petroleum geoscience and subsurface',
 'https://github.com/AlmazErmilov/FormationEval-an-Open-Benchmark-for-Oil-Gas-Geoscience-MCQ-Evaluation',
 'Repository CC BY 4.0; imported tracks have separate notices',
 'https://github.com/AlmazErmilov/FormationEval-an-Open-Benchmark-for-Oil-Gas-Geoscience-MCQ-Evaluation/blob/main/LICENSE',
 'Secondary regression benchmark only; excluded from training','Later upstream','Evaluation only','Technical Director','Use original track after rights review; inspect third-party notices for imported tracks; no refinery competence inference.')
source('D06','MIT OpenCourseWare','Educational materials','Chemical engineering fundamentals',
 'https://ocw.mit.edu/pages/privacy-and-terms-of-use/','CC BY-NC-SA 4.0; published AI-training policy prohibits commercial application',
 'https://ocw.mit.edu/pages/privacy-and-terms-of-use/','Excluded from commercial model training without separate permission','Exclude','Hold','Knowledge Engineer','Use alternative cleared sources; a course being free to read is not a commercial training grant.')

source('K01','Norman and Elizabeth Lieberman expertise','Expert knowledge candidate','Refinery equipment troubleshooting and field judgment',
 'https://www.lieberman-eng.com/personnel.htm','No Broadbridge engagement or training license; site states all rights reserved',
 'https://www.lieberman-eng.com/personnel.htm','Contracted original interviews and reviewed case authoring; licensed retrieval and SFT','Launch','Permission required','Technical Director','Scope new cases with expert and independent reviewer; establish ownership and former-client clearance before capture.')
source('K02','Lieberman published equipment and troubleshooting works','Published proprietary content','Distillation vacuum systems heat transfer pumps and controls',
 'https://www.sciencedirect.com/book/monograph/9780128161616/understanding-process-equipment-for-operators-and-engineers',
 'Publisher-managed publication; commercial AI rights not established','https://www.lieberman-eng.com/',
 'Licensed retrieval; training only if separately granted','Launch','Permission required','Managing Director','Map each title edition coauthor publisher and permitted use; author consent alone may not clear published content.')
source('K03','Henry Kister Distillation Diagnostics','Published proprietary content and expert candidate','Distillation diagnostics',
 'https://onlinelibrary.wiley.com/doi/book/10.1002/9781119640165','Publisher page lists 2025 AIChE copyright; no Broadbridge license',
 'https://onlinelibrary.wiley.com/doi/book/10.1002/9781119640165',
 'Potential licensed diagnostic references and separately contracted expert review','Launch alternate','Permission required','Technical Director','Assess access and reviewer availability; contract rights holder separately from any personal engagement.')
source('K04','PetroSkills Campbell Gas Course','Published proprietary content','Gas conditioning treatment dehydration and processing',
 'https://www.petroskills.com/en/training/courses/gas-conditioning-and-processing-g-4~p2995',
 'Commercial course and books; no AI reuse grant verified','https://www.petroskills.com/en/training/courses/gas-conditioning-and-processing-g-4~p2995',
 'Potential licensed material and specialist instruction for later gas practice','Later gas','Permission required','Technical Director','Seek institutional rights and qualified instructors only with a paid gas-practice mandate.')
source('K05','Customer incidents procedures historian and maintenance records','Prospective private data','Site-specific equipment and operating evidence',
 '','No customer data grant recorded','',
 'Tenant-only retrieval by default; shared SFT only with explicit permission','Launch','Permission required','Managing Director','Secure data agreement covering site use de-identification derived cases training and retention; keep underlying source restrictions.')
source('K06','OEM manuals standards and customer licensed references','Prospective proprietary reference','Equipment limits curves standards and operating context',
 '','Item-specific owner and license must be identified','',
 'Permission-controlled retrieval; shared training held by default','Launch support','Permission required','Managing Director','Inventory exact editions and licenses; customer possession does not establish redistribution or model-training rights.')

R={r['source_id']:r for r in records}
assert len(R)==len(records)==32
d=base('Public data models and expert knowledge for SLM training')
d.paragraphs[0].text='Broadbridge Oil and Gas training sources'
d.paragraphs[0].style='Title'
d.paragraphs[2].text='Broadbridge4096 subsidiary program  |  Research reviewed 23 September 2026'
d.paragraphs[2].runs[0].font.size=Pt(9)
d.core_properties.title='Broadbridge Oil and Gas training sources'
md=['# Broadbridge Oil and Gas training sources','Public data models and expert knowledge for SLM training','Research reviewed 23 September 2026','']

def hyperlink(p,label,url):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    e=OxmlElement('w:hyperlink'); e.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); pr=OxmlElement('w:rPr'); color=OxmlElement('w:color'); color.set(qn('w:val'),'18364D'); pr.append(color); r.append(pr)
    t=OxmlElement('w:t'); t.text=label; r.append(t); e.append(r); p._p.append(e)

def rich(p,text):
    for part in re.split(r'(\[(?:P|M|D|K)\d{2}\])',text):
        key=part[1:-1]
        if key in R and R[key]['source_url']: hyperlink(p,part,R[key]['source_url'])
        else: p.add_run(part)

def mdlinks(text):
    for key,r in R.items():
        if r['source_url']: text=text.replace(f'[{key}]',f'[{key}]({r["source_url"]})')
    return text

def p(text):
    rich(d.add_paragraph(),text); md.extend([mdlinks(text),''])

def h(text,level=1,new=False):
    if new: d.add_page_break()
    d.add_heading(text,level); md.extend(['#'*(level+1)+' '+text,''])

def table(headers,rows,widths,size=10.5):
    t=word_table(d,headers,rows,widths,size)
    for row in t.rows[1:]:
        for c in row.cells:
            text=c.text
            if re.search(r'\[(?:P|M|D|K)\d{2}\]',text):
                q=c.paragraphs[0]; q.clear(); rich(q,text)
                for run in q.runs: run.font.size=Pt(size)
    md.append('| '+' | '.join(headers)+' |'); md.append('|'+'|'.join(['---']*len(headers))+'|')
    md.extend('| '+' | '.join(mdlinks(str(v)).replace('\n','; ') for v in row)+' |' for row in rows); md.append('')

h('1 Recommended source strategy')
p('Build Broadbridge Process SLM around a general open-weight model, a small body of licensed expert cases, and a versioned technical reference library. Public material supplies fundamentals and evidence; experts supply the diagnostic distinctions that make those fundamentals useful at a plant. This sourcing program supports Broadbridge Oil & Gas, a subsidiary of Broadbridge4096, and the existing refining-first development plan.')
p('Start with distillation and vacuum troubleshooting. Steam, heat transfer, pumps and instrumentation provide supporting context. Broad petroleum, chemistry and geoscience collections should enter only when they address a measured gap in that scope. Model size, download counts and dataset row counts do not demonstrate refinery competence.')
table(['Input layer','What to acquire','How it contributes'],[
 ('Foundation model','Qwen3-8B and a 4B challenger; an alternative from Mistral [M01] [M02] [M03]','Starting weights for supervised fine-tuning; general language and tool-use behavior.'),
 ('Public engineering references','Selected DOE sourcebooks, EPA process descriptions and CSB investigations [P01] [P02] [P04] [P05]','Citable retrieval and reviewed fundamentals, evidence and escalation examples.'),
 ('Licensed expert knowledge','Original case capture with experts such as Norman Lieberman; published works require separate rights [K01] [K02]','Diagnostic questions, competing explanations, counterexamples and boundaries.'),
 ('Engineering tools','CoolProp properties and selected IDAES models [P07] [P08]','Checkable calculations and simulated examples with declared assumptions.'),
 ('Customer evidence','Explicitly permitted site records [K05]','Tenant-specific retrieval; optional shared examples only when the grant permits it.')],[1.35,2.8,3.05],10.5)
h('Initial acquisition priorities',2)
p('Prioritize one lead process expert, one independent reviewer, three DOE sourcebook collections, one carefully reconstructed CSB case, a small licensed reference library and two model candidates. Preserve the 2,000-example training target. Do not begin with a bulk collection of millions of petroleum-tagged rows.')
p('This brief and the accompanying 32-entry source register are a sourcing assessment. The listed materials have not been ingested, licensed on Broadbridge’s behalf, or approved for production training. Candidate experts have not been contacted or appointed. Source IDs in brackets link to the reviewed public pages; the register records rights evidence and the next acquisition action.')

h('2 Public sources for the refining launch',new=True)
table(['Source','Useful knowledge and proposed artifact','Admission condition'],[
 ('DOE Steam Systems [P01]','Steam generation, distribution, condensate and heat-transfer context. Produce short cited explanations and missing-data questions.','Check each document and third-party contribution; retain edition.'),
 ('DOE Process Heating [P02]','Heat balances, heating efficiency and system-performance vocabulary. Produce reviewed fundamental examples.','Use cleared passages; an efficiency guide is not a site operating procedure.'),
 ('DOE Pumping Systems [P03]','Pump/system relationships and performance evidence. Add supporting diagnostic questions.','Review contributed material, including industry-association content.'),
 ('CSB refinery investigations [P04]','Event sequences, evidence gaps and escalation failures. Create historical cases with an explicit decision time.','Review report attachments and rights; keep final findings out of pre-event prompts.'),
 ('EPA AP 42 Chapter 5 [P05]','Refinery process taxonomy and emissions context. Build a versioned glossary and reference collection.','Record the applicable section and date; validate intended use of factors.'),
 ('CoolProp [P07]','Property calculations for approved fluids and ranges. Produce tool calls with units, outputs and validation records.','MIT software license; check backends, ranges and property-model suitability.'),
 ('IDAES [P08]','Balances and constrained process simulations. Generate controlled perturbations and tool exercises.','BSD-style terms; audit dependencies and validate the model. Simulations are not plant observations.'),
 ('EIA Open Data [P06]','Refinery throughput, capacity and energy context. Use a dated API/tool feed.','Preserve series definitions and revisions; aggregate statistics are not troubleshooting labels.')],[1.7,3.25,2.25],10)
p('DOE distinguishes government information from protected contributed content. Its policy supports selecting and attributing cleared material, not declaring every hosted PDF reusable. EPA likewise notes document-specific copyright conditions. The source register links these policies. This review identifies collections and intended uses; acquisition must inspect the actual files and passages.')
h('Additional data with narrower uses',2)
p('NIST Chemistry WebBook is valuable for property validation, but Standard Reference Data has a distinct copyright regime; bulk reuse needs its own review [P09]. Volve is a later upstream resource under custom terms, including restrictions on selling the licensed material [P10]. NETL EDX is a discovery portal whose individual datasets require separate selection and rights checks [P11]. Neither is a substitute for refining cases.')

h('3 Existing models and their roles',new=True)
p('Use existing models in four distinct ways: starting weights, a drafting teacher, passage embeddings and passage ranking. A language model is not a verified source of engineering facts. Preserve the selected model’s license, revision, tokenizer, chat template and training configuration separately from the rights ledger for data.')
table(['Exact Hugging Face repository','Role and recommendation','Check before selection'],[
 ('Qwen/Qwen3-8B [M01]','Primary 8B-class baseline for the planned adapter. Official card lists Apache 2.0.','Measure refining performance before and after tuning with identical retrieval and tools.'),
 ('Qwen/Qwen3.5-4B [M02]','Smaller challenger; official card lists Apache 2.0. Use text-only scope initially.','Validate hybrid architecture, adapter targets and serving compatibility; vision capability is outside v1.'),
 ('mistralai/Ministral-3-8B-Instruct-2512 [M03]','Alternative, with an 8.4B language model and vision encoder. Apache 2.0.','Card describes FP8 instruction weights; establish a supported training format and reload test.'),
 ('Qwen/Qwen3-32B [M04]','Optional locally hosted teacher and larger-model comparison. Apache 2.0.','Draft only from cleared evidence. An independent expert accepts the answer; teacher output is not ground truth.'),
 ('Qwen/Qwen3-Embedding-0.6B [M05]','Embedding baseline for semantic retrieval. Apache 2.0.','Compare with lexical search and a hybrid combination using engineer-labeled relevant passages.'),
 ('Qwen/Qwen3-Reranker-0.6B [M06]','Rank retrieved passages by relevance. Apache 2.0.','Measure retrieval quality and latency separately from answer quality.')],[2.6,2.6,2.0],10.5)
h('How to use a teacher model',2)
p('Give the teacher an authorized evidence packet and an approved task template. Ask for a concise diagnostic brief, missing-data questions or a tool-use example. Preserve supporting passage IDs and calculation outputs. Reject unsupported claims, copied protected passages and implausible variations; then obtain expert review. Synthetic examples inherit the permission constraints of their inputs.')
p('Maintain the existing QLoRA/SFT approach and compare it with the untuned model plus the same retrieval and tools. Treat model merging, broad continued pretraining and training from scratch as separate experiments requiring evidence of need. Do not send private expert or customer material to a hosted teacher unless the agreement and approved hosting controls permit that processing.')

h('4 Domain models and datasets that need qualification',new=True)
table(['Asset','Finding from the reviewed source','Broadbridge decision'],[
 ('PetroGPT model [M07]','Llama-3-Petro-Instruct-v1 labels itself Apache 2.0 but identifies a Llama 3 base.','Hold pending license-chain and data-provenance review; no demonstrated advantage for this refinery task.'),
 ('PetroGPT dataset [D01]','petro-dataset-v2 lists CC BY-NC 4.0; the displayed examples include broad chemistry topics.','Exclude from commercial training and teacher generation unless separate rights are secured.'),
 ('GeoGPT model [M08]','GeoGPT-R1-Preview is a 72B geoscience model under custom terms; its card specifies noncommercial research/education.','Hold commercial deployment and distillation use pending permission. It is outside the launch SLM size and domain.'),
 ('GeoGPT CoT QA [D02]','39,393 generated geoscience QA records; card lists CC BY 4.0 and source-publication license metadata.','Potential later subset after DOI, answer and relevance review. Dataset rights differ from model rights.'),
 ('PETRA [D03]','Retrieval/reranking collection; license field says other. Displayed source chunks include unrelated subject matter.','Hold bulk ingestion; inspect source permissions and relevant labels before any retrieval experiment.'),
 ('PetroBench [D04]','Hugging Face card lists MIT and a test split; the viewer shows schema/encoding issues in some material.','Secondary evaluation candidate after cleanup and technical review. Keep it out of training.'),
 ('FormationEval [D05]','Original 505-question petroleum-geoscience track; repository CC BY 4.0 with separate notices for imported tracks.','Use for later regression only after source review. MCQ performance does not validate refining diagnosis.')],[1.65,3.05,2.5],10)
p('ChemLLM addresses chemistry and molecular science. Its card distinguishes code licensing from model-weight terms and requests commercial-license contact; defer it for this launch [M09]. MIT OpenCourseWare’s published AI policy restricts commercial training, so exclude it without separate permission [D06].')
p('These dispositions are acquisition recommendations, not a model bake-off or a legal clearance. The reviewed PetroBench repository and a 2026 paper share a name; this brief does not assume they are the same release. No public benchmark should be presented as an unseen test of a foundation model whose pretraining data is unknown.')

h('5 Proprietary domains worth capturing',new=True)
p('The most valuable expert contribution is a record of what changed the diagnosis: the observation that ruled out a tempting explanation, the measurement that mattered, the limit of a rule of thumb, and the circumstances requiring escalation. Capture those distinctions in original, rights-cleared cases with observable outcomes and independent review.')
table(['Knowledge domain','Expert or content route','Training artifact'],[
 ('Distillation and vacuum systems','Norman Lieberman as an illustrative candidate; separately licensed published references [K01] [K02]','Competing diagnoses, evidence selection, misleading readings, operating-regime boundaries and counterexamples.'),
 ('Equipment and operator reasoning','Norman and Elizabeth Lieberman expertise; their coauthored and individually published work requires title-level review [K01] [K02]','Translation of field observations into testable engineering questions, with units and instrument context.'),
 ('Distillation diagnostics review','Henry Kister’s published diagnostic work and a qualified independent fractionation reviewer [K03]','Independent challenge cases, falsification questions and reviewer rubrics.'),
 ('Rotating equipment and controls','Contract qualified pump, compressor and instrumentation specialists; obtain exact OEM references [K06]','Equipment-specific evidence requirements; distinguish sensor faults from process behavior.'),
 ('Gas conditioning and processing','Potential institutional licensing and specialist instruction through PetroSkills/Campbell materials [K04]','Later dehydration, treatment and processing curricula after practice approval.'),
 ('Actual site incidents','Customer operators, process engineers and reliability teams under a data agreement [K05]','Timestamped observations, approved interventions, outcomes and unresolved uncertainties.')],[1.8,2.9,2.5],10.5)
h('Expert acquisition package',2)
p('Propose a lead-expert engagement for original interviews and case development, plus a separately assigned reviewer. Plan an initial 8–12 recorded sessions of 60–90 minutes, supported by preparation and follow-up. This is a proposed workload, not an appointment or a guarantee of case yield. Secure recording and transcription rights before capture.')
p('License original interviews and published material separately. Specify retrieval, fine-tuning, synthetic derivatives, customer display, academy reuse, sublicensing and post-termination model use. Confirm publisher, coauthor, employer and former-client interests. Use retainers and scoped licensing; any equity remains a Broadbridge4096 decision. Name and likeness permissions do not imply expert endorsement of generated answers.')

h('6 Training curriculum and source mixture',new=True)
p('Use the existing target of 2,000 approved SFT examples as a planning constraint. The proposed mixture below allocates each record to its principal source; all contributing sources remain in its lineage. It is a curation target, not a claim that these examples already exist or an empirically optimal training ratio.')
mix=[('Original expert cases',600,250,0,50,100,1000),('Cleared public sources',180,80,0,150,90,500),('Validated simulations and tools',0,0,300,0,0,300),('Explicitly licensed customer cases',120,70,0,0,10,200)]
assert [sum(r[i] for r in mix) for i in range(1,7)]==[900,400,300,200,200,2000]
table(['Principal source','Diagnosis','Missing data','Tool use','Grounding','Escalation','Total'],mix+[('Total',900,400,300,200,200,2000)],[2.1,.85,.9,.75,.9,.9,.8],10)
p('This is 50% original expert material, 25% cleared public material, 15% validated calculations/simulations and 10% customer cases. If shared customer rights are unavailable, replace that 200-record allocation with cleared original expert cases while preserving the task totals. Teacher-generated variants remain in their source category and are separately tagged as synthetic; they do not count as new independent incidents.')
table(['Proposed topic emphasis','Examples','Scope'],[
 ('Distillation',700,'Launch diagnostic workflow'),
 ('Vacuum systems',500,'Launch diagnostic workflow'),
 ('Heat transfer and steam',400,'Supporting equipment and calculations'),
 ('Pumps and hydraulic context',200,'Supporting evidence within launch cases'),
 ('Instrumentation and evidence quality',200,'Sensor context, uncertainty and escalation'),
 ('Total',2000,'No upstream or drilling capability implied')],[2.5,.9,3.8],10.5)
h('Separate source families before generating examples',2)
p('Retain the program target of 300 incident/scenario families: 180 training, 60 development and 60 locked test. Derive the 2,000 training examples only from approved training families. Group incident retellings, book chapters, translations, paraphrases and simulation variants before splitting. Split simulation families by meaningful regime or topology, not adjacent random parameter rows.')
p('The existing acceptance suite remains 120 tasks from 60 locked cases plus 180 challenge prompts. Public historical cases are useful regression material but may already be known to base models. Use fresh, confidential, expert-reviewed cases to assess generalization; label synthetic cases and public cases separately in results.')

h('7 Turn sources into auditable training records',new=True)
p('Use a controlled transformation pipeline: select a business task, clear its sources, preserve the originals, extract evidence, group related material, assign a split, author examples, verify engineering content and publish a versioned release. Retrieval indexing follows the same rights and split controls as training.')
table(['Stage','Required result','Accountable owner'],[
 ('Source selection and capture','Document/version ID, original file hash, rights evidence, exact page or timestamp and capture date.','Knowledge Engineer'),
 ('Technical interpretation','Units, stream/equipment identity, measurement reliability, boundary conditions and observed versus inferred facts.','Technical Director'),
 ('Case authoring','Decision-time prompt, approved evidence, answer rubric and permitted tool calls; source lineage retained.','Knowledge Engineer'),
 ('Calculation validation','Reproducible tool inputs, software/model version, range checks and independently checked result.','Process Applications Engineer'),
 ('Independent review','Reviewer findings, technical corrections and acceptance record; reviewer did not author the item.','Technical Director'),
 ('Dataset release','Frozen manifest, source permissions, family split, deduplication report and tested export.','Head of Product & AI')],[1.55,3.85,1.8],10.5)
h('Illustrative case structure',2)
p('Example topic: a vacuum column’s reported absolute pressure has risen while separation performance has deteriorated. This is an original outline for authoring, not a validated plant diagnosis or an account attributed to Lieberman.')
p('The prompt includes only available observations and their timestamps. The target response identifies what can and cannot be concluded, asks for gauge basis and instrument verification, and organizes the evidence needed to distinguish process-load, condensation and vacuum-system explanations. It requests approved calculations where needed and states the boundary for specialist review. It does not invent a setpoint or issue a control action.')
p('Attach an expert-authored rubric listing plausible hypotheses, discriminating observations, unsupported leaps and escalation conditions. Keep the actual resolution and post-event findings in a separately protected answer record. Derive a diagnostic task, a missing-data task and a counterexample within the same family and split. Only accepted records enter the training release.')
p('Train concise explanations tied to observable evidence and tool results. Preserve disagreement and unresolved cases as uncertainty examples; do not force every incident into a confident single-cause answer. Remove duplicated passages, OCR mistakes, contradictory units and embedded document instructions before approval.')

h('8 Rights provenance and evaluation controls',new=True)
p('Use the register as an acquisition queue. Candidate means worth evaluating; Item review means the exact file or row needs review; Permission required means no grant is recorded; Hold excludes the material from the commercial pipeline; Evaluation only excludes it from training. Every entry currently says Not ingested. None of these labels records a production approval.')
table(['Permission or asset','Record and control'],[
 ('Expert background material','Actual owner, title/edition, author and publisher interests; signed grant and permitted uses.'),
 ('New interviews and cases','Recording, transcription, editing, model training, derivative examples, attribution and post-termination terms.'),
 ('Customer and OEM content','Tenant, site and authorized purpose; separate grants for shared training, external hosting, display and redistribution.'),
 ('Model weights and code','Base and derivative licenses, exact revisions, notices and deployment/distribution obligations.'),
 ('Dataset lineage','Source IDs and hashes, page/timestamp, family/split, transformations, synthetic flag and teacher/tool versions.'),
 ('Review and lifecycle','Author, independent reviewer, acceptance status, expiry, revocation, deletion obligations and affected releases.')],[2.15,5.05],10.5)
p('The Managing Director owns licensing commitments within delegation, supported by parent legal. Broadbridge4096 decides material IP transactions. The Knowledge Engineer maintains evidence; the Technical Director accepts engineering content; the Head of Product & AI enforces approved use in dataset and product releases. A free download, a book purchase or a customer’s access does not by itself record all of these permissions.')
h('Controls that preserve a meaningful test',2)
p('Exclude held-out answer keys and resolutions from training, teacher prompts, retrieval indexes, tool responses and caches. An evaluation may retrieve authorized evidence that would have been available at the question’s decision time. This requires explicit corpus filtering; keeping test rows out of an SFT file alone is insufficient.')
p('Audit source overlap with public benchmarks and evaluate base-model contamination where possible. Use blinded expert review, independent calculations and source-grounding checks. Report performance by topic, source class and real versus synthetic case. Keep private test answers outside general staff and customer retrieval.')
p('Maintain separate grants for retrieval and weight training. Prefer retrieval for frequently changing or revocable content. Removing a source from storage does not prove its influence has been removed from trained weights; any training contract must address retention, post-termination rights and the handling of affected model releases.')

h('9 Eight week sourcing and qualification program',new=True)
table(['Timing','Concrete deliverable','Accountable owner'],[
 ('Weeks 1–2','Confirm distillation/vacuum task coverage; rank the 32 sources; choose expert and reviewer candidates; define rights and family-split rules.','Technical Director'),
 ('Weeks 1–3','Negotiate only the required expert, publisher and pilot-data grants; record each permitted use before acquisition.','Managing Director'),
 ('Weeks 2–4','Prepare a proposed starter set of 20–30 selected documents or authorized excerpts; audit provenance, extraction and relevance.','Knowledge Engineer'),
 ('Weeks 3–6','Capture and review original expert cases; reconstruct permitted public cases; validate one property tool and one small simulation.','Technical Director'),
 ('Weeks 5–7','Run untuned model and retrieval comparisons using separate development material; measure supporting-passage retrieval and diagnostic quality.','Head of Product & AI'),
 ('Week 8','Target 50–60 cleared training families, a separate development sample and up to 500 approved training examples; provide rights, review and cost evidence for the next funding decision.','Head of Product & AI')],[.95,4.55,1.7],10.5)
p('These are the sourcing work packages within the existing eight-week first tranche, not additional budget or new employees. The seven launch roles, contracted experts and parent shared services remain as planned. The 2,000-example/300-family target remains a later program milestone, not an eight-week promise.')
h('Source admission scorecard',2)
p('First require valid rights for the proposed use, traceable provenance and exclusion of confidential material without permission. Then rank eligible sources on proposed weights: task relevance 35%, technical reliability 25%, case/evidence richness 20%, extraction quality 10% and acquisition effort 10%. These weights are management hypotheses. No high relevance score overrides missing rights.')
p('Proceed with adapter training when cleared case diversity, engineering acceptance and the untuned baseline are sufficient to test a measurable improvement. Continue source capture when gaps remain; do not fill the quota with unverified synthetic answers. Require a paid customer mandate and new specialists before adding gas-processing or upstream corpora to a commercial practice.')
h('Deliverables to carry into the training build',2)
p('Maintain the source register, signed grants, case schema, topic/task matrix, approved source library, separated case families, engineer-labeled retrieval queries, tool validation records and a dataset release manifest. The first model experiment should be reproducible from those records. Public-source findings and model-license observations in this brief are dated 23 September 2026 and must be rechecked for the chosen revisions.')
p('The companion CSV and JSON register contain the direct source links, rights-evidence links, priority, disposition, owner and next action for all 32 entries. The existing SLM Development Plan remains the schedule, staffing, budget and release-gate reference.')

(OUT/'docx').mkdir(parents=True,exist_ok=True)
(OUT/'training-sources').mkdir(parents=True,exist_ok=True)
d.save(OUT/'docx'/'Broadbridge_Oil_and_Gas_Training_Source_Synthesis.docx')
(ROOT/'Broadbridge-Oil-and-Gas-Training-Source-Synthesis.md').write_text('\n'.join(md),encoding='utf-8')
with (OUT/'training-sources'/'Broadbridge_Training_Source_Register.csv').open('w',encoding='utf-8-sig',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(records[0])); writer.writeheader(); writer.writerows(records)
(OUT/'training-sources'/'Broadbridge_Training_Source_Register.json').write_text(json.dumps(records,indent=2,ensure_ascii=False),encoding='utf-8')
print(json.dumps({'sources':len(records),'sft_examples':sum(r[-1] for r in mix),'docx':str(OUT/'docx'/'Broadbridge_Oil_and_Gas_Training_Source_Synthesis.docx')}))
