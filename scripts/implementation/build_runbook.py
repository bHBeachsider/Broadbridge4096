from pathlib import Path
import sys,json,math,re
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from build_oil_gas_documents import base,table
from plan_data import TASKS,REFS
from aws_start import pages as aws_pages,AWS_REFS,build_kit
FINAL_DAY=max(t['finish_day'] for t in TASKS)
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'output/multimodal-implementation'
PAGES=[]
def page(title,*blocks):PAGES.append((title,blocks))
def p(s):return ('p',s)
def h(s):return ('h',s)
def steps(*s):return ('steps',s)
def tab(headers,rows,widths=None):return ('table',(headers,rows,widths))
def code(s):return ('code',s)

page('1 Implementation instructions and schedule',
p('This runbook supplies the missing implementation layer beneath the CTO plan. The companion workbook contains70 tasks with inputs, instructions, outputs, acceptance checks, dependencies, accountable owners, reviewers, effort and an editable weekly Gantt. The format-examples ZIP demonstrates the file contracts on original synthetic data. Source acquisition, production parsing, GPU training and deployment remain implementation work.'),
p(f'The earlier24-week schedule was a top-level target. The AWS-first task order and daily effort limits now produce a proposed{FINAL_DAY}-working-day baseline, about{math.ceil(FINAL_DAY/5)} weeks. Gate dates are calculated from the workbook inputs. B01-B07 assume an already authorized AWS sandbox and run alongside source/business mobilization. G0 controls admitted source work. Rebaseline the24-week budget and first-eight-week funding limit before execution.'),
tab(['Decision','Calculated baseline','Required evidence'],[[t['id'],f"Day{t['finish_day']} / week{math.ceil(t['finish_day']/5)}",t['title']] for t in TASKS if t['phase']=='Gate'],[.7,1.7,4.8]),
steps('Open Start here. Enter an approved kickoff date; leave it blank to use working-day offsets. Weeks are five working days; weekends are excluded and holidays are not yet modeled.','Use Tasks for status, estimates and dependencies. Task detail contains the actual procedure and acceptance evidence. Use Source routing and Processing recipes to decide how each existing lead is handled.','Daily leveling assumes up to8 hours per named role and24 pooled expert/reviewer hours. These are peak reservation assumptions, not new hires. Source authors and independent reviewers must be separate people.','Review Capacity after changes. The formulas propagate dependencies but do not automatically re-level resources after edits. Reassign or move tasks if weekly demand exceeds confirmed capacity.','Compress the schedule only with an already-cleared corpus, earlier case production, validated tools, or additional explicitly funded capacity. Do not compress independent review by removing it.'))

page('2 Repository storage and toolchain',
p('Use a private Linux GPU environment for model work and an isolated CPU worker for source processing. Keep the original file, its derivatives and its permissions traceable. Source files are evidence; model weights are a separate software dependency. Never upload proprietary data to a public Hub repository.'),
tab(['Area','Contents and access'],[
['raw','Immutable originals, acquisition evidence and asset hashes. Intake identity writes; ordinary users cannot browse.'],
['extracted','Parser JSON, OCR text, page images, tables and figure regions. Still inherits source restrictions.'],
['curated','Reviewed evidence, case packets, split manifests, approved train/dev JSONL and image assets.'],
['retrieval','Authorized chunks, lexical/vector indexes and citation mapping. Access groups are enforced outside the model.'],
['models and runs','Pinned base snapshots; dependency locks; configs; adapters; logs and evaluation IDs. No raw mailbox or full client archive.'],
['evaluation','Separate evaluator prompts and gold answers. Trainer, teacher and production retrieval identities have no answer access.']],[1.3,5.9]),
h('Implementation choices to prove in the baseline'),
p('Use Python workers and a versioned object store with a PostgreSQL metadata catalog. Docling is the extraction candidate for PDFs and office files; test OCR/layout quality on actual engineering pages. Use Python email parsers for EML/MBOX, PyArrow for typed measurements, and an approved transcription engine for recordings. Use PostgreSQL full-text plus a vector index initially; retain exact embedding revision and dimensions.'),
p('Model tooling: PyTorch, Transformers, TRL, PEFT, Datasets, Accelerate and Pillow, with CUDA/driver compatibility verified together. GPU sizing follows a real processed batch, not parameter count alone. Separate training and serving dependency locks. Model-card examples may use development branches; select and pin a working commit instead of leaving main floating.'),
p('TasksB01-B07 create these environments. Record container digest, OS, GPU, driver, CUDA, package versions, model and processor hashes, and successful train/save/reload results in compatibility_report.json. A successful inference call alone does not prove training compatibility.'))

page('3 Source admission and the existing documents',
p('Start from the51 source leads already identified:32 general entries and19 Norm Lieberman entries. The Source routing sheet preserves their IDs and URLs. It is a disposition register, not a claim that files have been acquired or training permission obtained.'),
steps('Inventory the exact item: author, coauthor, title, edition/date, publisher, original URL, current custodian and any former-client material. Use separate assets for attachments and third-party figures.','Record permitted acquisition, local storage, OCR, retrieval, training, evaluation, derivative weights, internal users, processors, retention and revocation. Link the actual grant or terms evidence. Public availability is not an admission flag.','Assign allowed routes. A retrieval-only grant permits authorized retrieval; it does not permit exporting that text into SFT. Model licenses, software licenses and source-content rights are separate records.','Obtain the authoritative full asset through permitted access. Log acquisition time and checksum. If only a catalog, preview or search snippet is available, keep an acquisition lead; do not create technical answer labels from missing pages.','Choose the recipe in the workbook, admit the allowed uses and retain denied uses. Use an explicit state transition signed by the responsible owner; never let an ingestion script infer legal permission.'),
tab(['Existing lead','Concrete implementation route'],[
['N02-N08 and N18 OGJ articles','Obtain cleared full article; apply R01; create evidence blocks and reviewed cases only for granted uses.'],
['N09-N11 N13-N17 N19','Acquire full authoritative paper/chapter/article first. Catalogs, previews and unverified issue boundaries stay R11.'],
['N12 video library and K01 expertise','Separate grants for recordings, frames, quizzes and original interviews. Transcript/frame route R07; quizzes chosen for tests remain isolated.'],
['P01-P05 public references','Review the exact document and third-party credits; use R01/R02 after clearance. P04 incident questions separate pre-event evidence from hindsight.'],
['M01-M06 models and P07-P08 tools','Download model snapshots via R09; build checked software tools via R10. Neither is an article-ingestion task.'],
['D01-D06 and M07-M09','Preserve existing holds, exclusions or later/evaluation-only dispositions. No automatic bulk download or commercial-use inference.']],[2.0,5.2]),
p('Norm Lieberman remains an illustrative expert candidate. Original contracted cases are a preferred acquisition route; published works may require publisher and coauthor permissions. The plan assumes no appointment or granted rights.'))

page('4 Intake state machine and repeatable processing',
tab(['State','Entry action','Output and exit condition'],[
['Discovered','Register a lead and its rights questions.','source_id created; no content ingestion approval.'],
['Acquired in quarantine','Authorized copy; hash, MIME check and isolated malware scan.','asset_id and original bytes; safe to parse under approved use.'],
['Admitted','Verify grant, intended route and access boundary.','Signed use flags; unknown or expired grants stay held.'],
['Extracted','Run the recipe on an immutable original.','Text/tables/images plus parser revision and output hashes.'],
['Normalized','Check order, symbols, metadata, units and duplicates.','Canonical evidence with source locators and family ID.'],
['Reviewed','Technical reviewer accepts facts/labels and limitations.','Case/evidence review state; rejected regions remain visible.'],
['Released','Freeze split, permitted-use manifest and dataset/index revision.','Only approved route-specific artifacts reach retrieval or trainer.']],[1.3,2.5,3.4]),
steps('Assign IDs deterministically: source item, asset hash, page/region, case family, chunk and example. Store source URLs as metadata, never use a changing URL as the only identity.','For every run store batch_id, input SHA256, parser/version, configuration hash, timestamp, rights version, output paths/hashes, success/failure and reviewer decisions.','Make each worker idempotent on input hash plus parser/config/rights revision. Retry failed or changed items. Never append duplicates because a job was restarted.','Place parsing failures, unreadable pages, missing figures and expired grants in a rejection queue with reason and owner. A zero-length extraction is a failure, not a successful empty document.','When a grant changes, find every descendant chunk, image, example, index, run and adapter. Delete/restrict retrieval copies and caches. A trained parameter cannot be assumed to forget a source; retire or retrain the affected adapter when required.'),
p('TasksC01 and C12 build the intake handler and audit controls. Required negative checks include a duplicate file, a renamed duplicate, a malicious document, a missing attachment, a revoked source and an interrupted batch restart.'))

page('5 PDFs HTML and scanned pages',
h('Digital source recipe R01'),
steps('Keep the original PDF or permitted HTML snapshot. Record edition, full article boundaries and acquisition rights. HTML extraction removes menus and repeated site navigation, not authorship, qualifications or reference notes.','Extract native text and layout first. Save a canonical JSON representation containing ordered blocks, page number, bounding box, tables, headings and figure captions; save Markdown only as a readable review view.','Export figures and page renders as PNG where text and lines matter. Link each image to page, original coordinates, caption and source ID. Preserve the full page for audit even when a training example uses a crop.','Represent tables as cells with row/column indices, headers, merged-cell interpretation, units and footnotes. Keep a rendered table image as evidence; do not flatten a complex table into an ambiguous stream.','Compare every page count and all critical numeric regions with the original. For the first batch inspect every page; after stable extraction, keep complete automated checks plus risk-based human page sampling and100% review of critical values used in labels.'),
h('Scan recipe R02'),
steps('Render pages near300dpi as a starting acquisition setting. Preserve original pixels; orientation/deskew/contrast changes create derivatives with recorded transforms. Do not discard margins containing units, legends or revisions.','Run OCR in the approved environment. Retain region coordinates and OCR quality flags. Handwriting, equations, small tag text and line crossings require human verification; OCR confidence alone cannot approve engineering content.','Normalize to a documented coordinate convention: page number1-based; bounding boxes in original-image pixels with top-left origin; retain width/height and each crop-to-original transform.','Flag uncertain characters and missing information explicitly. Do not let an LLM silently repair a pressure value, exponent, minus sign or flow arrow. Preserve extracted_raw and reviewed_text separately.'),
p('Outputs: documents.jsonl, regions.jsonl, page PNGs, figure PNGs, structured table JSON and an extraction audit. A PDF is not directly supplied to the SFT trainer. Reviewed passages and images later become examples or retrieval evidence. TasksC03-C04 and C10 implement these steps.'))

page('6 Work papers emails and recorded expertise',
h('Office files and work papers'),
p('For DOCX/PPTX, extract paragraphs, tables, slide notes, drawings and image relationships. Save a fixed-layout preview for review and preserve the native original. Decide explicitly whether tracked changes, comments, speaker notes and hidden slides are included. Equations retain a source image and a reviewed transcription with variables and units. Calculations must be independently recomputed before they become answer labels. Implement inC05.'),
h('Email threads'),
steps('Use an authorized EML/MBOX export. For MSG/PST, use an approved export tool first; do not assume a document parser can safely read the mailbox container.','Parse MIME structure, Message-ID, In-Reply-To, References, From/To, timestamps and attachments. Preserve timezone and thread order. Resolve linked documents only under separate authorized acquisition.','Replace repeated quoted bodies with references to earlier messages while retaining the original. Remove signatures and personal details from the training derivative, preserving necessary professional role context.','Create a case packet with facts known at the question time and the later diagnosis/outcome in distinct fields. A later email answer must never appear in the diagnostic prompt.','Review permissions for each correspondent, client details and attachment. Export redacted thread JSONL and evidence spans; create teaching records only after technical and rights review. Implement inC07.'),
h('Interviews audio and video'),
steps('Agree the capture brief and recording/reuse permissions before recording. Ask for symptoms, context, alternatives, decisive observations, failed approaches and limits of the lesson.','Transcribe with time offsets and speaker IDs using an approved processor. Have the expert correct names, equipment tags, technical terms and numerical statements against the audio.','Export segments containing start/end time, speaker, text and source ID. For video, select instructional frames at known times; preserve diagrams at usable resolution and link to the related transcript.','Create text or image/text cases from the approved transcript and frames. Raw audio/video is retained as evidence; it is not an input modality for release-one SFT. Implement inC09.'),
p('Keep the expert account and independent reviewer assessment distinct. If the expert disagrees with a label, preserve and adjudicate the disagreement; do not turn it into a supposedly settled answer.'))

page('7 Diagrams photographs and measurement data',
h('Visual annotation recipe R05'),
steps('Retain the native drawing and an approved image export. Record drawing revision, equipment/system, orientation and confidentiality. Treat CAD objects as separate structured data; do not assume their attributes are visible in a raster image.','Select an overview plus at most two relevant crops as the initial annotation policy. Keep tags, legends and connecting line segments inside the selected regions. The processor pixel/token budget may require fewer or smaller crops.','Annotate visible equipment tags, values/units, arrow direction, line/edge endpoints, region coordinates and legibility. Label a crossing without a clear junction as ambiguous. Mark unreadable text rather than guessing it.','Pair the view with a specific task: locate a tag, extract a stated value, identify a bounded connection, compare evidence or ask for missing information. Avoid teaching unrestricted whole-plant connectivity from one sheet.','Have an independent engineer compare labels with the original. Wrong diagram rotations, mirrored images and arbitrary geometric augmentation can change meaning and are excluded.'),
h('Numeric recipe R06'),
steps('For CSV/XLSX/historian exports, preserve values, source timestamps, timezone, sampling interval, sensor quality, units and basis. Export formulas and calculated values separately when spreadsheets contain both.','Use typed Parquet columns for bulk measurements; use CSV for small review exchanges. Record tag, event_time, value, unit, quality and source_revision. Missing readings are null, not zero.','Normalize units only with a documented reversible transform. Preserve gauge versus absolute pressure and actual versus standard flow conditions. Record assumptions and conversion version.','Align trends only with an explicit resampling/interpolation rule approved for the task. Preserve original series and report gaps. Calculate from numbers, not visually estimated pixels, when the numbers are available.','Generate a labeled PNG plot from the checked table. Link plot, numeric data and time window. Train visual observations on the image; route quantitative calculations to the checked numeric tool.'),
p('C06 and C08 produce reviewed visual_labels.jsonl, image assets, measurements.parquet and data_dictionary.json. Equipment photographs may support bounded observations, but this pilot does not certify damage severity, dimensions or fitness for service from pixels.'))

page('8 Canonical records and quality checks',
p('Use UTF-8 JSONL for reviewable record exchanges: one complete JSON object per line, no comments and no trailing commas. Use JSON for schemas/manifests, Parquet for large typed tables, and PNG/JPEG files for images. JSONL is the storage contract; a Python dataset plus decoded images is the trainer runtime contract.'),
tab(['Record','Required fields'],[
['Asset','asset_id; source_id; source hash; MIME; original path; grant_id; acquisition timestamp; edition; permitted uses; access groups; parent asset; derivative transform.'],
['Evidence block','block/chunk ID; original asset; family; page/region or time locator; raw and reviewed text; table/image IDs; units; revision; review state.'],
['Case family','family_id; parent incident/publication/scenario; source assets; topic/equipment; known-at-question-time facts; outcome; split; all derivative IDs.'],
['SFT example','example_id; family_id; split; source/evidence/grant IDs; review provenance; image_paths; prompt messages; completion messages; schema version.'],
['Evaluation task','task_id; family; prompt evidence; expected evidence/answer in evaluator-only store; tolerances; critical-error rubric; subtype; author/reviewer.'],
['Model run','run_id; model SHA; processor/config hashes; dependency lock; data snapshot; target modules; seed; learning parameters; GPU/time/cost; checkpoints.']],[1.3,5.9]),
h('Admission checks before any export'),
p('Verify unique IDs, allowed route, nonexpired grant, asset existence and hashes, schema types, timestamps and units. Compare row/column context and equation variables to originals. Preserve source-to-derivative lineage. Reject a reviewed answer with no supporting evidence or an image label outside its parent region.'),
p('Use exact file hashes for identical assets and normalized-text/image similarity to identify candidate near duplicates. An engineer resolves similarity clusters. Republishing the same incident, cropping its diagram or generating ten paraphrases does not create ten independent case families.'),
p('Schema validity is necessary but does not prove technical truth. Every production SFT answer needs an identifiable independent review. A release contains accepted and rejected counts, reasons and coverage statistics. C10-C12 and D09 implement these checks.'))

page('9 Case construction splits and leakage control',
steps('Create the family before generating questions. Group an incident, its article, email follow-ups, diagrams, revisions, crops and synthetic variants under one family. Group a parameterized synthetic scenario by shared template/underlying event as appropriate.','Assign180 families to training,60 to development and60 to locked testing. Keep at least90 usable visual families among the180 training families and30 among each60-family dev/test partition. If sources cannot support that inventory, change the scope or schedule; do not inflate paraphrases.','For each case, record operating regime, what was observed, what was unknown, alternatives, discriminating checks, outcome and limitations. Make question-time evidence explicit; keep later outcome evidence out of the prompt.','Generate teaching variants only after the family split. Use existing approved evidence and concise checked explanations. Optional model-generated drafts retain teacher/version/prompt/source metadata and require expert review.','Build2,000 training examples from training families only. Development and test prompts are additional records. Use development data for configuration choices; open locked tests only after freezing the candidate.','Keep test-family answer material out of SFT, teacher prompts, retrieval indexes, tool caches and general developer logs. Evaluator questions may contain authorized evidence needed to solve the task, but the gold answer remains isolated.','Run exact and near-duplicate checks over all exported surfaces. If a locked test reveals a defect, fix it and use fresh held-out cases for the affected generalization claim; preserve the old test only as a regression check.'),
tab(['Training curriculum','Text','Image and text','Total'],[['Diagnostic work',625,275,900],['Missing evidence questions',300,100,400],['Checked tools and calculations',225,75,300],['Source grounding',170,30,200],['Escalation and scope',180,20,200],['Total',1500,500,2000]],[3.1,1.0,1.7,1.4]),
p('D01-D09 implement the case inventory, split manifests, authoring and audit. Keep case-family counts and example counts separate in reports. Public benchmark performance is supplementary and cannot establish refining competence.'))

page('10 Exact supervised training file format',
p('The authoritative record is canonical JSONL containing prompt, completion and image_paths plus governance fields. A prompt is the question with the information available to answer it. A completion is the independently reviewed target response. The loader sends only model-input fields to the trainer.'),
code('''{"schema_version":"bb-sft-1","example_id":"EX-0001",
 "family_id":"F-001","split":"train",
 "purpose":"production_candidate",
 "source_ids":["SRC-001"],"evidence_ids":["EV-001"],
 "grant_ids":["GRANT-001"],
 "review":{"state":"accepted","author":"author-id",
   "reviewer":"independent-id","technical_approval":true},
 "image_paths":["images/figure-001.png"],
 "prompt":[
   {"role":"system","content":[{"type":"text",
     "text":"Use supplied evidence and cite its ID."}]},
   {"role":"user","content":[{"type":"image"},
     {"type":"text","text":"Describe trend EV-001 and its limits."}]}],
 "completion":[{"role":"assistant","content":[
   {"type":"text","text":"Reviewed answer with citation [EV-001]."}]}]}'''),
p('The block above is valid JSON formatted for reading. JSONL stores each complete object on one physical line; the ZIP contains actual parseable examples. The grant and people IDs above illustrate the fields and do not represent approvals.'),
steps('For text-only records, image_paths is[] and there is no image placeholder. Keep content blocks consistently typed. Do not add blank dummy images.','For visual records, the number and order of image placeholders must match image_paths. The ZIP loader opens each approved local image, converts it to RGB and returns images as a list of PIL image objects.','The runtime dataset has prompt, completion and images columns. Do not pass filenames as if the model can read them from text. Do not fetch arbitrary remote image URLs during training.','Preserve metadata in the canonical registry but strip it from the model input unless the task intentionally needs a field. Validation answers and rights documents never become prompt text.','Use the model-specific processor and chat template at batch time. Do not manually insert guessed image special tokens or pre-embed pictures into a text-only CSV.'),
p('For tool-call training, use the selected model’s actual tool-message schema plus a tools schema column and validated tool responses. Implement and smoke-test that adapter separately; the supplied minimal schema intentionally covers text/image brief training only.'))

page('11 Retrieval ingestion and engineering tools',
h('Document evidence goes into retrieval before training'),
steps('Read only admitted, reviewed evidence blocks. Split by semantic heading, starting with400-800 model-token chunks and about80 tokens of overlap where needed. These are tuning defaults, not required standards. Keep a table and its headers together or create explicit row groups.','Assign chunk_id and retain text, source/asset IDs, page/region, image IDs, family, revision, grant, access groups and split. A citation resolves through this mapping to the original.','Use the pinned embedding model to compute text vectors. Store vector dimension, model SHA, normalization and query/document formatting. Never mix vectors from different embedding revisions in one unversioned index.','Index approved text in a lexical index and approved vectors in a vector index. Filter user/project permissions, grant status and excluded test-family material before candidate retrieval, not merely after model generation.','Retrieve a candidate set, initially20 passages, then rerank authorized candidates and select a smaller context set within the model budget. The reranker does not receive unauthorized candidates.','Build a context pack with cited passages and approved image regions. Text embeddings alone do not provide image understanding; image assets reach the VLM through retrieved region links or user-selected authorized evidence.','Evaluate recall on engineer-labeled development queries. Start with a90% recall@20 target and inspect misses before changing chunking, embeddings or reranking. Rebuild/version the index when the embedding model changes.'),
h('Engineering calculations follow a separate route'),
p('P07 CoolProp and P08 IDAES are software candidates. Pin code, dependencies and any property backends; review the selected equations and validity ranges. Expose only named tools with typed inputs, explicit units, range checks and reproducible outputs. Independently validate a small reference set before creating synthetic examples from results.'),
p('Store tool name/version, input values/units, method, output values/units, validity flags and evidence in ToolResult JSON. A failed calculation is returned as an error with missing inputs, not replaced by model-generated numbers. Tool outputs may be supplied in training examples only after validation. E01-E06 implement this path.'),
p('Retrieval is read-time access to evidence; SFT changes model behavior. A searchable chunk is not automatically a good training example. Revocable, frequently changing or customer-specific facts generally remain in permission-controlled retrieval.'))

page('12 Hugging Face model acquisition and setup',
p('The primary candidate isQwen/Qwen3.5-4B; Qwen/Qwen3-8B is the text control. The embedding and reranking candidates areQwen/Qwen3-Embedding-0.6B andQwen/Qwen3-Reranker-0.6B. Ministral3 remains an alternate requiring its own format/training compatibility test. The selected model is downloaded as weights and configuration, not copied into the training-data folder.'),
steps('Review the exact model card, license and dependency requirements. Record repository ID, full commit SHA, purpose, approver and download date in model_manifest.json. Resolve the SHA from the repository; do not invent a version or leave main as the release revision.','Download original safetensors shards, shard index if present, config, generation config, tokenizer assets, chat template, processor/image-processor assets and notices. Keep all referenced assets needed to load. The supplied download_model.py requires APPROVED_MODEL_REVISION.','Verify hashes and shard completeness. Store model files privately under models. Reject unreviewed executable repository code; load with trust_remote_code=False when supported. A pickle-style checkpoint needs a separate security decision.','Create a compatible locked GPU environment. If the model requires unreleased Transformers/runtime code, use an approved exact commit and record the exception. Start with model-native precision; do not substitute an arbitrary GGUF or third-party4-bit inference build for training.','Load the processor and model from local paths. For the reference Qwen path, the documented class isQwen3_5ForConditionalGeneration with AutoProcessor. Test text, one image, two ordered images and a text-only row in a mixed batch.','Record GPU memory, processed lengths and response shape. Inspect language module names for LoRA targets. Run ten optimizer steps, save, reload and compare outputs before approving the compatibility matrix.'),
code('''# Reference download after approval and a pinned environment
export APPROVED_MODEL_REVISION=<full approved commit SHA>
python templates/download_model.py

# Local loading pattern for the selected Qwen model
processor = AutoProcessor.from_pretrained(
    base_dir, local_files_only=True, trust_remote_code=False)
model = Qwen3_5ForConditionalGeneration.from_pretrained(
    base_dir, local_files_only=True, trust_remote_code=False,
    torch_dtype=torch.bfloat16)'''),
p('These commands are implementation instructions for the project environment. No weights were downloaded or GPU execution performed while creating this plan. B02-B07 provide the required acceptance evidence.'))

page('13 From JSONL to a model training batch',
steps('Validate canonical records first: schema, allowed SFT use, review approval, unique example IDs, family partition, image hashes and local paths. Resolve evidence and grant IDs. Refuse test records on the training host.','Load JSONL as structured records. Decode every image to RGB and keep images in placeholder order. Build a Hugging Face Dataset containing prompt, completion and images. Keep the larger provenance record outside the trainer.','Apply the model’s own processor/chat template to representative text-only, one-image, two-image and mixed batches. The processor handles tokenization, image resizing/patching and required image tensors. Do not use a generic text tokenizer as a replacement.','Inspect the full processed sequence length including image expansion. The proposed initial admission cap is8192 tokens. Reject or restructure over-budget records before training; include an overview only if it fits. Record actual pixel policy and measured image-token counts.','Use prompt/completion training with completion-only loss. Verify the collator masks system/user context, image placeholders/tokens and padding while leaving the intended answer tokens supervised. Inspect actual tensors for this model revision.','For each audited batch, decode the supervised label positions and compare with the intended target answer. Check that no answer text was accidentally included in prompt evidence and no target answer disappeared behind a mask.','Use no packing initially. For VLM training, set trainer max_length=None only after enforcing the dataset-level cap; this prevents a later generic truncation from cutting required image tokens. Fail on oversized batches rather than silently trimming them.','Run a forward/backward pass. Require finite loss, expected trainable gradients and acceptable peak GPU memory. Then complete the ten-step save/reload test. Retain the batch audit as release evidence.'),
tab(['Before processor','After processor and collator'],[
['prompt and completion message objects','input_ids and attention_mask; model-specific conversation markers.'],
['images as ordered decoded image objects','Model-specific pixel tensors and grid/patch metadata.'],
['Reviewed target answer','labels with ignored positions masked to-100; target positions supervised.'],
['Asset/case metadata','Audit record outside the input tensors; not a hidden source of answers.']],[3.1,4.1]),
p('F03 and B06 own this boundary. Exact tensor keys depend on the model. The CPU fixture validator cannot verify token expansion, masks or GPU compatibility; the actual processor/collator tests are mandatory implementation tasks.'))

page('14 Fine tuning experiments and execution',
p('Begin with the existing vision-language model and LoRA adapters on audited language layers. Freeze the visual encoder and projector initially. List exact target modules from named_modules and verify that their names belong to the language stack. Avoid an indiscriminate all-linear target that also changes the vision system.'),
tab(['Initial proposal','Value or implementation rule'],[
['Data','Approved train release; dev used for selection; no test mount.'],
['Precision','BF16 when the GPU and model support it. QLoRA is a separate memory-saving experiment, not an automatic format conversion.'],
['Adapter','Rank16; alpha32; dropout0.05; no bias; explicit language module list.'],
['Optimization','One epoch; learning rate5e-5; batch1 per device; accumulation16; seed17. Effective batch equals batch × accumulation × GPU count.'],
['Sequence handling','Preflight cap8192 processed tokens; max_length=None after admission; packing=False; completion_only_loss=True.'],
['Experiment bounds','Compare ranks16/32 and LR5e-5/1e-4; at most eight development configurations per finalist; repeat the selected recipe with three seeds.'],
['Tracking','Base/processor/data/code hashes; config; trainable parameter count; loss; dev scores; GPU memory/hours; cost; checkpoint IDs.']],[1.5,5.7]),
steps('Run the baseline system with retrieval/tools before adaptation. Compare the text control and the same VLM with OCR-only versus image input to isolate the value of vision.','Use100/200/400-example family-grouped subsets for a pilot learning curve. Diagnose source, extraction, retrieval, label, perception and tool errors separately. More epochs cannot fix a wrong label.','Implement the reference recipe on the approved GPU environment. The ZIP template requires a separately created approved_training_config.json tied to saved approval/test evidence. It is not a substitute for B02-B07.','Stop on NaN/infinite loss, broken masks, image-count mismatch, leakage, unexpected trainable modules, exceeded spend or critical development regression. Preserve logs and classify the failure.','Train the selected final recipe on the frozen2,000-example release. Compare tuned and untuned versions using identical evidence/tools. Retain adaptation only if it yields the proposed five-point acceptance or15% reviewer-time improvement at matched quality.','Freeze the candidate and all dependencies before independent test evaluation. Do not choose a checkpoint using the locked test answers.'),
p('F04-F08 produce experiment configurations, reviewed development results, adapter checkpoints and a reproducible candidate. Learning rate, batch and image budget are starting proposals to validate, not guaranteed optimal settings.'))

page('15 Model export reload and deployment package',
steps('Save the PEFT adapter as adapter_model.safetensors with adapter_config.json. Preserve the exact base model revision and any additional saved modules. An adapter alone is not a complete standalone model.','Save the processor/tokenizer/chat template used for training with the release. Include dependency locks, module targets, dataset manifest, evaluation reports and approved generation settings.','Reload the base plus adapter in a fresh process and environment. Test deterministic text and visual fixtures, compare outputs/logits within a defined numerical tolerance and verify the intended adapter weights are active.','Use the proven Transformers path first. Adopt vLLM or another serving engine only after its exact architecture, multimodal processor and LoRA support pass compatibility tests. Model-card inference examples do not establish support for the trained adapter.','If deployment requires merging or quantization, create a new derived artifact with parent hashes. Confirm supported merge behavior and re-evaluate quality, image handling, memory and latency. Never label a quantized artifact equivalent without testing.','Serve through an authenticated internal application. Apply request limits, permission-filtered retrieval, allowed tool schemas and a validated DiagnosticBrief response. Reject unrecognized tool names and unauthorized evidence IDs server-side.','Version the entire system: base, adapter, processor, prompt, retrieval index, tool runtime, code and security configuration. A rollback restores a compatible package, not just a weight file.'),
tab(['Release file or record','Purpose'],[
['release_manifest.json','All component hashes and technical/security acceptance IDs.'],
['adapter files plus base reference','Learned parameter changes and the exact underlying model needed to use them.'],
['processor and dependency lock','Reproduce the image/text representation and runtime.'],
['dataset card and rights manifest','Training provenance, permitted use, exclusions and limitations.'],
['evaluation and operations records','Quality evidence, user instructions, monitoring, restore and rollback procedures.']],[2.4,4.8]),
p('H02-H03 implement the service and F08 validates the export. The Head of Product and AI owns release; Technical Director and parent security acceptance are prerequisites. No plant-control write path is included.'))

page('16 Independent evaluation implementation',
steps('Freeze the candidate and evaluation protocol before opening the locked set. Use120 tasks from60 independent held-out families:60 visual tasks from30 families and60 text tasks from30 families. Keep180 separate challenge prompts for visual, grounding/tool/scope and security behavior.','Separate prompt files from evaluator gold files. The inference service receives only task evidence and the question. Record prediction, citations, tool calls, latency and release ID in predictions.jsonl.','Have two independent qualified reviewers score material diagnostic cases. Reviewers do not certify their own authored work. Blind model identity and counterbalance task order to reduce familiarity effects in time comparisons.','Use fixed denominators and legibility labels defined before model execution. Unnecessary refusal on a sufficient-evidence task counts as failure. Refusing a preclassified legible field counts as a missed extraction; the model cannot exclude hard cases itself.','Score diagnostic usefulness, material-claim grounding, exact critical tag/unit/value extraction, bounded edge precision/recall, numeric correctness, appropriate escalation and critical failures. Store item-level decisions and adjudications.','Measure performance on the declared input workload and hardware, including image-token cap, five concurrent users, up to three tool calls and an800-token output target. Separate model latency from full reviewed-brief effort.','Report family-clustered paired comparisons and uncertainty, not just aggregate percentages. Inspect every modality and equipment slice. A finite zero-critical-error result does not establish that failures are impossible.','If locked testing triggers a material correction, obtain fresh held-out tasks for affected claims. Use exposed tasks for regression only. Defer the gate if new independent evidence or reviewer capacity is unavailable.'),
tab(['Proposed acceptance measure','Threshold'],[['Diagnostic tasks accepted without material correction','At least85%'],['Material factual claims supported by evidence or checked tools','At least95%; zero fabricated source IDs'],['Legible critical tags units values and numeric tool tasks','At least98%'],['Bounded connection precision and recall; correct escalation','At least95% for each'],['Critical errors unauthorized disclosure or actions','Zero observed in gate suite; no unresolved high security finding'],['Runtime and workflow','p95 brief within60seconds on declared workload; target25% less review effort']],[4.7,2.5]),
p('These are Broadbridge proposals, not statutory or industry-mandated thresholds. D08 and I01-I04 implement the evaluation and remediation records.'))

page('17 Security monitoring and operational handoff',
p('Build controls into every data transition and the application. A prompt telling the model to respect confidentiality is insufficient. Training data permissions and application access are enforced by identities, storage policies, retrieval filters and service code.'),
steps('Test user/project isolation across originals, extracted text, image URLs, retrieval candidates, reranking, caches, exports and logs. An image must inherit the same restrictions as its source text.','Test prompt injection embedded in a document, table, diagram and OCR text. Treat source content as data. Only the application can authorize tool calls, routes or grants.','Keep source/model downloads in a controlled acquisition job. Disable arbitrary browsing and shell execution in the engineering assistant. Apply resource limits and allowlisted tools. Scan packages and model artifacts; preserve software inventory.','Track request/release/evidence IDs, latency, input size, tool failures, rejected briefs and reviewer corrections. Limit sensitive prompt logging and define retention before use.','Quarantine user feedback. Reconfirm source rights, group it into the existing family, have an independent reviewer approve it and release through the normal dataset process. Never automatically retrain from accepted chat messages.','Exercise backup restore and rollback before pilot approval. Demonstrate the proposed four-business-hour recovery and24-hour maximum loss for review records, or revise those service targets explicitly.','Disable the affected capability for a credible critical engineering defect or disclosure. Notify the Technical Director or security lead; preserve restricted audit evidence; repair and repeat acceptance before restart.','Create one release record with technical acceptance, security acceptance and product release authorization. The MD controls operating commitments and funding; expert council influence does not override formal decision rights.'),
p('B01, E06, H04 and I02/I05/I06 supply access tests, revocation proof, monitoring, incident procedures and the release package. The production implementation must demonstrate each control on actual services; the format examples do not implement them.'),
h('Gate evidence to retain'),
p('Save access matrix and negative tests; dependency/container scans; malicious-document handling tests; prompt-injection results; validated tool tests; source revocation tests; data/adapter lineage; restore and rollback logs; named on-call contacts and user instructions.'))

page('18 Pilot execution and task completion',
steps('Before pilot access, name the internal users, available reviewers, accepted use cases and service window. Explain what constitutes an unsupported answer and how to inspect the original evidence.','Onboard a target of ten internal users using a supervised sample. Require evidence inspection, unit checking and acceptance/rejection before a brief enters engineering work.','Collect a target of100 reviewed briefs over the four-week pilot period including onboarding and closure. Record task complexity, accepted/rejected status, correction type, review time, cited evidence, model release and any escalation.','Second-review a weekly sample and all critical incidents. If the service generates a critical unsupported result, suspend the affected function; average usefulness cannot override the technical veto.','Compare observed effort with the manual baseline for comparable tasks. Count unsuccessful attempts and correction time. Report infrastructure plus review labor per accepted brief and a continuing service forecast.','Use pilot evidence to decide continue, narrow or stop. Future external customers, additional modalities or new practices require separate scope, data rights, evaluation and funding decisions.'),
h('A task is complete only when its output exists'),
p('The workbook status field starts at Not started. Change it only when the stated acceptance evidence exists and the designated reviewer has accepted it. The Gantt dates are estimates, not proof of completion. Link the evidence record in the task blocker/evidence field. Do not mark a task complete because its scheduled end date has passed.'),
h('Resource and cost interpretation'),
p('The detailed schedule contains owner effort and separate reviewer effort. The Capacity sheet compares estimated assigned hours with the earlier program-hour envelope and shows weekly peaks. Unassigned hours are not savings: they cover coordination, iterations, additional review, operations and uncertainty. Contractor estimates are capacity placeholders, not agreed appointments.'),
p('The prior4.25 FTE average and24-week budget do not automatically fund31 weeks. AtG0 reconcile salaries, protected time, expert availability, hosting duration and the first-eight-week spend limit. The current schedule assumes cleared initial assets are available at intake. Rights negotiations and unavailable experts can delay it further.'),
h('First practical handoff'),
p('AssignA01-A07, approve the environment and grants, then implementB01-B07 and the initialC/D/E tasks. Run the supplied CPU fixture demonstration to establish the data contract. Replace its synthetic records only with admitted, independently reviewed production candidates.'))

page('19 Worked format example and supplied files',
p('The ZIP includes a five-row synthetic pressure CSV, a plotted PNG, asset/evidence manifests, two actual JSONL teaching examples, a schema, a CPU validator and reference model-loader/training templates. There is no proprietary publication, customer information or claimed real-world diagnosis in the example.'),
('image',str(OUT/'format_examples/images/demo_pressure.png')),
steps('Follow raw/demo_measurements.csv into images/demo_pressure.png. Evidence manifests retain the original and derived hashes, source relationship and citation ID.','Read curated/train_demo.jsonl. One record asks about stated numeric values; the other asks for a bounded visual observation and what cannot be inferred. They share one family.','Run python validate_and_preview.py in an environment with Pillow and jsonschema. Expected output is two valid demo records, one image record and runtime columns prompt/completion/images.','Review templates/load_dataset.py to see local paths become actual image objects. Review download_model.py for the pinned snapshot acquisition step.','Review train_reference.py as an implementation starting recipe. It requires a separate approved configuration and real GPU compatibility, token, mask and reload tests. It has not been run against model weights here.'),
p('The example validator checks structure and file integrity. It cannot confer rights, independently certify technical content or validate model behavior. Do not count the fixture toward2,000 examples or300 families. The evaluation file is a format template, not an actual held-out test case.'))

page('20 Technical references and implementation boundaries',
p('Official documentation checked23 September2026. Exact package/model versions are intentionally selected through compatibility tasks rather than invented in advance. The specific file contracts, batch policies, schedules and acceptance checks in this runbook are Broadbridge implementation proposals.'),
('refs',None),
h('How to use the package'),
p('The workbook is the task and schedule authority. This runbook explains the formats and implementation procedures referenced by its section numbers. The ZIP gives concrete examples and limited reference code. The earlier CTO plan remains the source for business scope and governance except where this detailed schedule explicitly revises timing.'),
p('Production components still to build include the quarantine handler, per-format extractors, grant enforcement, annotation workflow, family/split manager, token preflight, evidence retrieval, engineering tools, evaluation runner and internal service. Their task-level acceptance checks are in Task detail. This delivery has not executed those production tasks.'),
p('Original research, prior plans, original source registers and the Energy Trading folder link are preserved. Trading materials remain organizational references.'))

def clean(s):
    if not isinstance(s,str):return s
    s=re.sub(r'([a-z])(?=\d)',r'\1 ',s)
    s=re.sub(r'(?<=\d)(?=[a-z])',' ',s)
    s=re.sub(r'(?<=[a-z])(?=[A-Z][0-9])',' ',s)
    s=re.sub(r'(?<=[a-z])(?=Qwen|AutoProcessor)',' ',s)
    s=re.sub(r'([:,])(?=\d{1,2}\b)',r'\1 ',s)
    s=re.sub(r'(\d) e([+-]\d)',r'\1e\2',s)
    return s.replace('Qwen 3','Qwen3').replace('p 95','p95').replace('is[]','is []').replace('to-100','to -100').replace('g 6 e.2 xlarge','g6e.2xlarge').replace('G6 e','G6e').replace('gp 3','gp3').replace('s 3://','s3://').replace('IMDSv 2','IMDSv2')
PAGES=aws_pages(p,h,steps,tab,code,TASKS)+PAGES
REFS=REFS+AWS_REFS
build_kit()
d=base('Multimodal SLM implementation runbook')
d.paragraphs[0].text='Broadbridge Oil and Gas';d.paragraphs[0].style=d.styles['Title']
d.paragraphs[2].text='AWS foundation model through domain training  |  24 September 2026'
d.styles['Normal'].font.size=Pt(10.5)
d.styles['Normal'].paragraph_format.space_after=Pt(6)
md=['# Broadbridge Oil and Gas multimodal SLM implementation runbook','']
for i,(title,blocks) in enumerate(PAGES):
    if i:d.add_page_break()
    d.add_heading(title,1);md.extend(['## '+title,''])
    for kind,value in blocks:
        if kind in ['p','h']:value=clean(value)
        if kind=='steps':value=[clean(s) for s in value]
        if kind=='p':d.add_paragraph(value);md.extend([value,''])
        elif kind=='h':d.add_heading(value,2);md.extend(['### '+value,''])
        elif kind=='steps':
            for n,s in enumerate(value,1):d.add_paragraph(f'{n}. {s}');md.append(f'{n}. {s}')
            md.append('')
        elif kind=='code':
            for line in value.splitlines():
                para=d.add_paragraph();para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=1
                run=para.add_run(line);run.font.name='Consolas';run.font.size=Pt(8.6)
            d.add_paragraph();md.extend(['```text',value,'```',''])
        elif kind=='table':
            headers,rows,widths=value
            rows=[[clean(v) for v in row] for row in rows]
            t=table(d,headers,rows,widths)
            for row in t.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after=Pt(3)
                        for r in para.runs:r.font.size=Pt(9.3)
            md.append('| '+' | '.join(headers)+' |');md.append('| '+' | '.join(['---']*len(headers))+' |')
            md.extend('| '+' | '.join(map(str,row))+' |' for row in rows);md.append('')
        elif kind=='image':d.add_picture(value,width=Inches(5.8));md.extend(['![Synthetic data plot](format_examples/images/demo_pressure.png)',''])
        elif kind=='refs':
            for id,label,url,note in REFS:
                para=d.add_paragraph();para.add_run(id+' '+label+' — '+note+' ')
                if url.startswith('https:'):
                    rel=para.part.relate_to(url,RT.HYPERLINK,is_external=True);hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),rel)
                    r=OxmlElement('w:r');rp=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'18364D');rp.append(col);r.append(rp);tx=OxmlElement('w:t');tx.text='Official source';r.append(tx);hl.append(r);para._p.append(hl)
                else:para.add_run(url)
                md.append(f'- {id} [{label}]({url}): {note}')
            md.append('')
d.save(OUT/'Broadbridge_AWS_Implementation_Runbook.docx')
(OUT/'Broadbridge_AWS_Implementation_Runbook.md').write_text('\n'.join(md),encoding='utf-8')
print(f'Saved {len(PAGES)}-section AWS-first implementation runbook')
